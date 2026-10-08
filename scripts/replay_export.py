"""Review an existing v0.4 export through current assembly, without SEC access.

This is regression evidence from a prior application output, NOT independent
verification against source filings. It does not make a permanent data archive.
"""
from argparse import ArgumentParser
from pathlib import Path
from openpyxl import load_workbook

from financial_historian.models import Company, Fact, Filing, CoverageRecord
from financial_historian.extract import assemble_bundle
from financial_historian.service import limit_statement_history
from financial_historian.export import build_workbook


def load_export(source: Path, company: Company):
    """Recover public source records, including v0.5+ contexts and dimensions."""
    workbook=load_workbook(source,read_only=True,data_only=True)
    facts=[]
    for row in workbook["Source Data"].iter_rows(min_row=5,values_only=True):
        if not row[1] or not isinstance(row[4],(int,float)):
            continue
        facts.append(Fact(tag=row[1],label=row[2],unit=row[3],value=row[4],start=row[5],end=row[6],
                          accession=row[7],filed=row[8],form=row[9],extraction=row[10],selection_reason=row[11] or "",
                          source_url=row[12] or "",statement=row[0],presentation_value=row[14],
                          source_instant=row[15],row_role=row[16] or "",section=row[17] or "",
                          context_id=(row[18] or "") if len(row)>18 else "",
                          dimensions=tuple(tuple(item.split("=",1)) for item in (row[19] or "").split("; ") if "=" in item) if len(row)>19 else (),
                          decimals=str(row[20]) if len(row)>20 and row[20] is not None else ""))
    filings=[]
    coverage=[]
    for row in workbook["Sources & Coverage"].iter_rows(min_row=5,values_only=True):
        if row[0]=="Filing":
            filings.append(Filing(row[1],row[2] or "10-K",row[4],row[3],"",row[7] or ""))
        elif row[0]=="Coverage":
            coverage.append(CoverageRecord(row[1],row[3],row[4],row[7] or "",row[5],row[6] or ""))
    workbook.close()
    bundle=limit_statement_history(assemble_bundle(company,filings,facts,coverage,[
        "REPLAY OF A PRIOR APPLICATION EXPORT. No live SEC request or independent source verification occurred. "
        "Source metadata is retained where present; records absent from the original export cannot be recovered here."
    ]))
    bundle.disclosures=[f for f in facts if f.statement=="Disclosure"]
    return bundle


def replay(source: Path, destination: Path, company: Company):
    if source.resolve()==destination.resolve():
        raise ValueError("Choose a different output path; preserve the original export.")
    bundle=load_export(source,company)
    destination.parent.mkdir(parents=True,exist_ok=True)
    destination.write_bytes(build_workbook(bundle))
    return bundle


if __name__=="__main__":
    parser=ArgumentParser(description=__doc__)
    parser.add_argument("source",type=Path)
    parser.add_argument("output",type=Path)
    parser.add_argument("--cik",required=True)
    parser.add_argument("--name",required=True)
    parser.add_argument("--ticker",default="")
    parser.add_argument("--sic",default="")
    args=parser.parse_args()
    bundle=replay(args.source,args.output,Company(args.cik,args.name,args.ticker,args.sic))
    print(f"Replayed {len(bundle.raw_facts)} prior source records into {len(bundle.periods)} annual columns.")
    print("Not independent live-filing verification.")
