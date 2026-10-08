"""Rebuild an existing export and verify every defined metric's Excel cache.

This checks consistency and regression behavior, not independent filing accuracy.
The source workbook is never modified. Replays are user-controlled artifacts,
not a permanent financial-data store.
"""
from argparse import ArgumentParser
from collections import Counter
from decimal import Decimal
from io import BytesIO
import json
import math
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from financial_historian.analysis import build_analysis, currency, DEFINITIONS
from financial_historian.export import build_workbook
from financial_historian.models import Company
from scripts.replay_export import load_export


def check(source, output, company):
    if source.resolve()==output.resolve():
        raise ValueError("Preserve the original workbook; choose a different output.")
    started=perf_counter()
    bundle=load_export(source,company)
    analysis=build_analysis(bundle)
    analyzed=perf_counter()
    payload=build_workbook(bundle,analysis=analysis)
    exported=perf_counter()
    workbook=load_workbook(BytesIO(payload),data_only=True,read_only=True)
    matrices={}
    for title in ("Historical Ratios","Return Bridge"):
        sheet=workbook[title]
        periods=next(sheet.iter_rows(min_row=4,max_row=4,values_only=True))[1:]
        for row in sheet.iter_rows(min_row=5,values_only=True):
            if row[0]:
                matrices[str(row[0]).split(" [")[0]]=dict(zip(periods,row[1:]))
    keys={d.key for d in DEFINITIONS}
    defined=[r for r in analysis.rows if r.definition.key in keys]
    checked=0
    for row in defined:
        cached=matrices[row.definition.label]
        for period,item in row.cells.items():
            actual=cached[period]
            if item.value is None:
                assert actual is None,(row.definition.key,period,"Expected explained blank")
            else:
                expected=item.value/(Decimal(1000000) if currency(item.unit) else 1)
                if item.unit=="pp":
                    expected*=100
                assert isinstance(actual,(float,int)) and math.isclose(actual,float(expected),rel_tol=1e-12,abs_tol=1e-10),(row.definition.key,period)
            checked+=1
    assert {"Per-share Basis","Filing Changes","Source Data","Sources & Coverage"}<=set(workbook.sheetnames)
    workbook.close()
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_bytes(payload)
    return {"ticker":company.ticker,"evidence":"replay of supplied export; not independent live verification",
            "periods":len(analysis.periods),"defined_measures":len(defined),"cached_cells_checked":checked,
            "measures_with_any_value":sum(any(c.value is not None for c in r.cells.values()) for r in defined),
            "populated_metric_cells":sum(c.value is not None for r in defined for c in r.cells.values()),
            "basis_reviews":len(analysis.basis_reviews),"validation":dict(Counter(c.status for c in analysis.checks)),
            "load_and_analysis_seconds":round(analyzed-started,3),"export_seconds":round(exported-analyzed,3),
            "metric_period_counts":{r.definition.key:sum(c.value is not None for c in r.cells.values()) for r in defined},
            "latest_values":{r.definition.key:str(r.cells[analysis.periods[-1]].value) for r in defined}}


if __name__=="__main__":
    parser=ArgumentParser(description=__doc__)
    parser.add_argument("source",type=Path)
    parser.add_argument("output",type=Path)
    parser.add_argument("--cik",required=True)
    parser.add_argument("--name",required=True)
    parser.add_argument("--ticker",required=True)
    parser.add_argument("--sic",required=True)
    args=parser.parse_args()
    print(json.dumps(check(args.source,args.output,Company(args.cik,args.name,args.ticker,args.sic)),indent=2))
