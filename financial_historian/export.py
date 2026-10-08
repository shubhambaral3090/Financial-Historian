"""Build a sourced workbook without converting reported units or inventing values."""
from __future__ import annotations

from collections.abc import Callable
from datetime import date
from io import BytesIO
import math
import re
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Fact, STATEMENTS, StatementBundle, StatementRow
from .normalization import concept_for
from .presentation import fact_rejection


MAX_ROWS = 1_048_576
MAX_COLUMNS = 16_384
MAX_TEXT_LENGTH = 32_767
MAX_PRESENTATION_PERIODS = 20
HEADER_ROW = 4
FIRST_DATA_ROW = 5

# Only an entire currency unit is scaled. USD/shares, shares, pure and other
# compound units remain in their reported unit. No currency conversion occurs.
CURRENCIES = frozenset("""
AED AFN ALL AMD ANG AOA ARS AUD AWG AZN BAM BBD BDT BGN BHD BIF BMD BND
BOB BOV BRL BSD BTN BWP BYN BZD CAD CDF CHE CHF CHW CLF CLP CNY COP COU
CRC CUC CUP CVE CZK DJF DKK DOP DZD EGP ERN ETB EUR FJD FKP GBP GEL GHS
GIP GMD GNF GTQ GYD HKD HNL HRK HTG HUF IDR ILS INR IQD IRR ISK JMD
JOD JPY KES KGS KHR KMF KPW KRW KWD KYD KZT LAK LBP LKR LRD LSL LYD
MAD MDL MGA MKD MMK MNT MOP MRU MUR MVR MWK MXN MXV MYR MZN NAD NGN
NIO NOK NPR NZD OMR PAB PEN PGK PHP PKR PLN PYG QAR RON RSD RUB RWF
SAR SBD SCR SDG SEK SGD SHP SLE SLL SOS SRD SSP STN SVC SYP SZL THB
TJS TMT TND TOP TRY TTD TWD TZS UAH UGX USD USN UYI UYU UYW UZS VED
VES VND VUV WST XAF XCD XOF XPF YER ZAR ZMW ZWL
""".split())

_ILLEGAL_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_INK = "152535"
_TEAL = "007F82"
_PALE = "EFF6F8"
_MUTED = "526878"
_LINE = "DCE6EB"

# Reuse immutable style objects across the potentially large raw fact sheet.
_BODY_FONT = Font(name="Aptos", size=10, color=_INK)
_BOLD_FONT = Font(name="Aptos", size=10, color=_INK, bold=True)
_MUTED_FONT = Font(name="Aptos", size=10, color=_MUTED)
_HEADER_FONT = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=_INK)
_PALE_FILL = PatternFill("solid", fgColor=_PALE)
_TOTAL_FILL = PatternFill("solid", fgColor="E8F1EE")
_BODY_BORDER = Border(bottom=Side(style="hair", color=_LINE))
_TOTAL_BORDER = Border(top=Side(style="thin", color="8DA59D"), bottom=Side(style="thin", color=_LINE))
_TEXT_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
_VALUE_ALIGNMENT = Alignment(horizontal="right", vertical="center")
_PLAIN_ALIGNMENT = Alignment(vertical="top")
_NUMBER_FORMAT = '#,##0.00;[Red](#,##0.00);0.00'


def _write_text(cell, value: object) -> None:
    """Explicit string cells prevent Excel formula injection, including '='."""
    value = "" if value is None else str(value)
    if len(value) > MAX_TEXT_LENGTH:
        raise ValueError("A workbook text field exceeds Excel's 32,767-character limit.")
    if _ILLEGAL_TEXT.search(value):
        raise ValueError("A workbook text field contains characters Excel cannot store.")
    cell.value = value
    cell.data_type = "s"


def _safe_source_url(value: str) -> str | None:
    """Only official SEC HTTPS data and archive paths become clickable links."""
    if not isinstance(value, str) or re.search(r"[\s\\\x00-\x1f]", value):
        return None
    try:
        parsed = urlsplit(value)
        if (parsed.scheme != "https" or parsed.hostname not in
                {"sec.gov", "www.sec.gov", "data.sec.gov"} or
                parsed.username or parsed.password or parsed.port not in {None, 443}):
            return None
        if "%" in parsed.path or ".." in parsed.path or "//" in parsed.path:
            return None
        if not parsed.path.startswith(("/Archives/edgar/data/", "/api/xbrl/", "/submissions/")):
            return None
    except ValueError:
        return None
    return value


def _source_link(cell, source: str) -> None:
    safe = _safe_source_url(source)
    if safe:
        cell.hyperlink = safe


def _display_unit(unit: str) -> str:
    return f"{unit} million" if unit in CURRENCIES else unit


def _valid_fact(row: StatementRow, period: str, fact: Fact, statement: str) -> str | None:
    return fact_rejection(row, period, fact, statement)


def _fact_comment(fact: Fact, display_unit: str, rejection: str | None = None) -> Comment:
    text = "\n".join([
        f"{'Generated row identifier (no SEC tag)' if fact.tag.startswith('html:') else 'SEC tag'}: {fact.tag}", f"Reported label: {fact.label}",
        f"Source/XBRL value (unscaled): {fact.value}",
        f"Statement presentation value (unscaled): {fact.presentation_value if fact.presentation_value is not None else 'Not supplied; source value used'}",
        f"Reported unit: {fact.unit}",
        f"Display unit: {display_unit}",
        f"Annual display period start: {fact.start or 'Instant / not provided'}",
        f"Annual display period end: {fact.end}",
        f"Original instant context: {fact.source_instant or (fact.end if fact.start is None else 'Not an instant fact')}",
        f"Row role: {fact.row_role or 'Reported statement row'}",
        f"Reported section: {fact.section or 'Not supplied'}",
        f"Filed: {fact.filed}", f"Form: {fact.form}",
        f"Accession: {fact.accession}", f"Source filing: {fact.source_url}",
        f"Extraction: {fact.extraction}",
        f"Selection reason: {fact.selection_reason or 'Not supplied by extraction'}",
        "A later comparative figure is not necessarily a confirmed restatement.",
        *( [f"Not displayed: {rejection}"] if rejection else [] ),
    ])
    if len(text) > MAX_TEXT_LENGTH or _ILLEGAL_TEXT.search(text):
        raise ValueError("A provenance comment exceeds Excel's text limits or contains invalid characters.")
    comment = Comment(text, "Fishtorian")
    comment.width = 460
    comment.height = 340
    return comment


def _base_sheet(workbook: Workbook, title: str, company: str, note: str, columns: int):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    _write_text(sheet.cell(1, 1), f"{company} | {title}")
    sheet.cell(1, 1).font = Font(name="Aptos", size=18, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=_INK)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, columns))
    sheet.row_dimensions[1].height = 34
    _write_text(sheet.cell(2, 1), note)
    sheet.cell(2, 1).font = Font(name="Aptos", size=10, color=_MUTED)
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, columns))
    sheet.row_dimensions[2].height = 30
    return sheet


def _headers(sheet, headers: list[str], row_number: int = HEADER_ROW) -> None:
    for col, value in enumerate(headers, 1):
        cell = sheet.cell(row_number, col)
        _write_text(cell, value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _TEXT_ALIGNMENT
    sheet.row_dimensions[row_number].height = 30


def _style_row(sheet, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.font = _BODY_FONT
        cell.alignment = _TEXT_ALIGNMENT if col in {1, 2} else _PLAIN_ALIGNMENT
        cell.border = _BODY_BORDER
        if row % 2:
            cell.fill = _PALE_FILL


def _unavailable_periods(rows: list[StatementRow], periods: list[str]) -> list[str]:
    """Keep every selected column, while explicitly recording entirely blank ones."""
    occupied = {period for row in rows for period in row.values}
    return [period for period in periods if period not in occupied]


def _cash_summary_rank(row: StatementRow) -> int | None:
    ranks = {
        "cash_begin": 0, "restricted_cash_begin": 0,
        "cfo": 1, "cfo_continuing": 1,
        "cfi": 2, "cfi_continuing": 2,
        "cff": 3, "cff_continuing": 3,
        "exchange_effect": 4, "restricted_exchange_effect": 4,
        "cash_change": 5, "restricted_cash_change": 5,
        "cash_end": 6, "restricted_cash_end": 6,
    }
    keys = {concept.key for fact in row.values.values() if (concept := concept_for(fact))}
    ranked = {ranks[key] for key in keys if key in ranks}
    return next(iter(ranked)) if len(ranked) == 1 else None


def _statement_note(rows: list[StatementRow]) -> tuple[str, set[str]]:
    currencies = {row.unit for row in rows if row.unit in CURRENCIES}
    unit_note = f"{next(iter(currencies))} million" if len(currencies) == 1 else "Currency amounts in millions; currency shown in each row"
    return f"{unit_note} · Other units shown in row labels · Blank = unavailable · Reported annual figures", currencies


def _statement_label(row: StatementRow, currencies: set[str]) -> str:
    if row.unit in CURRENCIES and len(currencies) == 1:
        return row.label
    return f"{row.label} ({_display_unit(row.unit)})"


def _write_statement(workbook, title: str, statement: str, company: str,
                     rows: list[StatementRow], periods: list[str], displayed: dict,
                     summary: bool = False) -> int:
    columns = max(2, len(periods) + 1)
    has_sections = not summary and any(row.section for row in rows)
    # Unclassified historical rows must not appear to belong to the preceding
    # section (for example an old asset line below the equity section).
    sections = [(row.section or "Other reported line items") if has_sections else "" for row in rows]
    section_count = sum(bool(section) and (index == 0 or section != sections[index - 1])
                        for index, section in enumerate(sections))
    _check_dimensions(len(rows) + section_count + HEADER_ROW, columns, title)
    note, currencies = _statement_note(rows)
    sheet = _base_sheet(workbook, title, company, note, columns)
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.tabColor = _TEAL if title in STATEMENTS else _MUTED
    sheet.freeze_panes = "B5"
    sheet.column_dimensions["A"].width = 52
    for col in range(2, columns + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 16
    guidance = ("Cash Flow Detail contains every reported cash-flow row. Values here are reported totals, not calculated estimates."
                if summary else "Select a value for its SEC filing; its cell note records the original label, tag, unit and filing date.")
    if not rows:
        guidance = ("No supported statement values are available for this cash summary. Cash Flow Detail may contain individual rows; review Sources & Coverage."
                    if summary else "No supported statement values were extracted. Review Sources & Coverage for filing availability and gaps.")
    _write_text(sheet.cell(3, 1), guidance)
    sheet.cell(3, 1).font = _MUTED_FONT
    sheet.cell(3, 1).alignment = _TEXT_ALIGNMENT
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)
    sheet.row_dimensions[3].height = 28
    _headers(sheet, ["Line item", *periods])
    withheld, row_number = _write_statement_rows(
        sheet, statement, rows, periods, displayed, FIRST_DATA_ROW, currencies,
        sections=sections, summary=summary,
    )
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:{get_column_letter(columns)}{max(HEADER_ROW, row_number - 1)}"
    return withheld


def _write_statement_rows(sheet, statement: str, rows: list[StatementRow],
                          periods: list[str], displayed: dict, row_number: int,
                          currencies: set[str], *, sections: list[str],
                          summary: bool = False) -> tuple[int, int]:
    """Use one cell writer for the statement tabs and the analyst matrix."""
    columns = max(2, len(periods) + 1)
    previous_section = ""
    withheld = 0
    for row, section in zip(rows, sections):
        if section and section != previous_section:
            for col in range(1, columns + 1):
                cell = sheet.cell(row_number, col)
                cell.fill = _PALE_FILL
                cell.font = _BOLD_FONT
            _write_text(sheet.cell(row_number, 1), section)
            sheet.row_dimensions[row_number].height = 25
            row_number += 1
        previous_section = section
        label = _statement_label(row, currencies)
        _write_text(sheet.cell(row_number, 1), label)
        sheet.row_dimensions[row_number].height = max(22, 14 * math.ceil(len(label) / 49))
        for col in range(1, columns + 1):
            cell = sheet.cell(row_number, col)
            cell.font = _BOLD_FONT if row.is_total or summary else _BODY_FONT
            cell.alignment = _TEXT_ALIGNMENT if col == 1 else _VALUE_ALIGNMENT
            if row.is_total or summary:
                cell.fill = _TOTAL_FILL
                cell.border = _TOTAL_BORDER
        for col, period in enumerate(periods, 2):
            fact = row.values.get(period)
            if fact is None:
                continue
            cell = sheet.cell(row_number, col)
            rejection = _valid_fact(row, period, fact, statement)
            cell.comment = _fact_comment(fact, _display_unit(row.unit), rejection)
            if rejection:
                withheld += 1
                continue
            presented = fact.presentation_value if fact.presentation_value is not None else fact.value
            cell.value = presented / 1_000_000 if row.unit in CURRENCIES else presented
            cell.number_format = "#,##0;[Red](#,##0);0" if row.unit == "shares" else _NUMBER_FORMAT
            _source_link(cell, fact.source_url)
            displayed.setdefault(fact, []).append(f"'{sheet.title}'!{cell.coordinate}")
        row_number += 1
    return withheld, row_number


def _write_data_sheet(workbook: Workbook, bundle: StatementBundle,
                      periods: list[str], displayed: dict) -> int:
    """Reliance-style annual matrix; technical observations live in Source Data.

    Column A is the reported line item, and B through at most U are chronological
    fiscal period ends. The same columns are repeated for each statement block.
    All selected periods stay visible here, including empty statement edges.
    """
    columns = max(2, len(periods) + 1)
    all_rows = [row for statement in STATEMENTS for row in bundle.statements.get(statement, [])]
    unit_note, currencies = _statement_note(all_rows)
    sheet = _base_sheet(workbook, "Data Sheet", bundle.company.name,
                        f"Latest 20 fiscal periods maximum · {unit_note}", columns)
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.tabColor = _TEAL
    sheet.freeze_panes = "B6"
    sheet.column_dimensions["A"].width = 52
    for col in range(2, columns + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 16
    _write_text(sheet.cell(3, 1), "Selected annual statements in one matrix. Select a value for its SEC source and cell note. Full original records: Source Data.")
    sheet.cell(3, 1).font = _MUTED_FONT
    sheet.cell(3, 1).alignment = _TEXT_ALIGNMENT
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)
    sheet.row_dimensions[3].height = 28
    row_number = HEADER_ROW
    withheld = 0
    for statement in STATEMENTS:
        rows = bundle.statements.get(statement, [])
        for col in range(1, columns + 1):
            cell = sheet.cell(row_number, col)
            cell.fill = PatternFill("solid", fgColor=_TEAL)
            cell.font = _HEADER_FONT
        _write_text(sheet.cell(row_number, 1), statement)
        sheet.row_dimensions[row_number].height = 26
        row_number += 1
        _headers(sheet, ["Reported line item", *periods], row_number)
        row_number += 1
        has_sections = any(row.section for row in rows)
        sections = [(row.section or "Other reported line items") if has_sections else "" for row in rows]
        section_count = sum(bool(section) and (index == 0 or section != sections[index - 1])
                            for index, section in enumerate(sections))
        _check_dimensions(row_number + len(rows) + section_count, columns, "Data Sheet")
        if not rows:
            _write_text(sheet.cell(row_number, 1), "No supported values — see Sources & Coverage")
            sheet.cell(row_number, 1).font = _MUTED_FONT
            sheet.cell(row_number, 1).alignment = _TEXT_ALIGNMENT
            sheet.row_dimensions[row_number].height = 30
            row_number += 1
        else:
            block_withheld, row_number = _write_statement_rows(
                sheet, statement, rows, periods, displayed, row_number, currencies,
                sections=sections,
            )
            withheld += block_withheld
        row_number += 1
    # A later page may contain Balance Sheet or Cash Flow; do not repeat the
    # first block's "Profit & Loss" band above a different statement.
    sheet.print_title_rows = "1:3"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:{get_column_letter(columns)}{row_number - 2}"
    return withheld


def _check_dimensions(rows: int, columns: int, sheet_name: str) -> None:
    if rows > MAX_ROWS or columns > MAX_COLUMNS:
        raise ValueError(f"{sheet_name} exceeds Excel's {MAX_ROWS:,}-row / "
                         f"{MAX_COLUMNS:,}-column limit. No data was silently dropped.")


def build_workbook(bundle: StatementBundle, progress: Callable[[str], None] | None = None, *, analysis=None) -> bytes:
    """Return an XLSX with reported statement rows and cell-level provenance.

    Periods are chronological and must already have the application's 20-period
    presentation limit applied. Monetary units display in millions; source values
    are unscaled in Source Data. Inconsistent facts
    remain visible there but are withheld from the misleading statement cell.
    """
    report = progress or (lambda message: None)
    report("Checking workbook values and source records…")
    if any(statement not in STATEMENTS for statement in bundle.statements):
        raise ValueError("An unsupported statement worksheet was supplied; no rows were dropped.")
    periods = sorted(set(bundle.periods))
    if len(periods) > MAX_PRESENTATION_PERIODS:
        raise ValueError("Apply the latest 20 fiscal-period selection before building the workbook; "
                         "export must match the online statement selection.")
    try:
        for period in periods:
            if date.fromisoformat(period).isoformat() != period:
                raise ValueError
    except (TypeError, ValueError):
        raise ValueError("Annual period labels must be ISO dates (YYYY-MM-DD).") from None

    raw_facts = list(bundle.raw_facts) + [f for f in bundle.disclosures if f not in bundle.raw_facts]
    seen = set(raw_facts)
    for rows in bundle.statements.values():
        for row in rows:
            for period, fact in row.values.items():
                if period not in periods:
                    raise ValueError(f"A statement contains period {period} absent from the bundle; no data was dropped.")
                if fact not in seen:
                    raw_facts.append(fact)
                    seen.add(fact)
    for fact in raw_facts:
        if isinstance(fact.value, bool) or not isinstance(fact.value, (int, float)) or not math.isfinite(fact.value):
            raise ValueError("Financial facts must contain finite numeric values; missing values belong in blank cells.")
        if fact.presentation_value is not None and (
                isinstance(fact.presentation_value, bool) or
                not isinstance(fact.presentation_value, (int, float)) or
                not math.isfinite(fact.presentation_value)):
            raise ValueError("Statement presentation values must contain finite numeric values when supplied.")
    _check_dimensions(len(raw_facts) + HEADER_ROW, 21, "Source Data")
    _check_dimensions(len(bundle.filings) + len(bundle.coverage) + len(bundle.warnings) + HEADER_ROW,
                      8, "Sources & Coverage")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Fishtorian annual statements"
    workbook.properties.creator = "Fishtorian"
    displayed: dict[Fact, list[str]] = {}
    withheld = 0
    unavailable_periods: dict[str, list[str]] = {}
    for statement in STATEMENTS:
        report(f"Formatting {statement.lower()}…")
        rows = bundle.statements.get(statement, [])
        unavailable_periods[statement] = _unavailable_periods(rows, periods)
        if statement == "Cash Flow":
            ranked = [(rank, index, row) for index, row in enumerate(rows)
                      if (rank := _cash_summary_rank(row)) is not None]
            summary_rows = [row for _, _, row in sorted(ranked)]
            withheld += _write_statement(workbook, statement, statement, bundle.company.name,
                                         summary_rows, periods, displayed, summary=True)
        else:
            withheld += _write_statement(workbook, statement, statement, bundle.company.name,
                                         rows, periods, displayed)
    cash_rows = bundle.statements.get("Cash Flow", [])
    withheld += _write_statement(workbook, "Cash Flow Detail", "Cash Flow", bundle.company.name,
                                 cash_rows, periods, displayed)
    report("Formatting the analyst Data Sheet…")
    withheld += _write_data_sheet(workbook, bundle, periods, displayed)

    report(f"Writing {len(raw_facts):,} source observations…")
    data = _base_sheet(workbook, "Source Data", bundle.company.name,
                       "Raw source values and optional statement presentation values retain their reported units. Original instant contexts are separate from annual display periods. No monetary scaling or currency conversion is applied on this sheet.", 21)
    _headers(data, ["Statement", "SEC tag / row identifier", "Reported label", "Reported unit", "Source/XBRL value (unscaled)",
                    "Annual display period start", "Annual display period end", "Accession", "Filed", "Form", "Extraction", "Selection reason",
                    "Source filing", "Displayed cells", "Statement presentation value (unscaled)",
                    "Original source instant", "Row role", "Reported section", "Context", "Dimensions", "Decimals"])
    for row_number, fact in enumerate(raw_facts, FIRST_DATA_ROW):
        _style_row(data, row_number, 21)
        values = [fact.statement, fact.tag, fact.label, fact.unit, fact.value, fact.start, fact.end,
                  fact.accession, fact.filed, fact.form, fact.extraction, fact.selection_reason,
                  fact.source_url, "; ".join(displayed.get(fact, [])), fact.presentation_value,
                  fact.source_instant, fact.row_role, fact.section, fact.context_id,
                  "; ".join(f"{axis}={member}" for axis,member in fact.dimensions), fact.decimals]
        for col, value in enumerate(values, 1):
            cell = data.cell(row_number, col)
            if col in {5, 15} and value is not None:
                cell.value = value
                cell.number_format = "#,##0.##########;[Red](#,##0.##########);0"
            elif value is not None:
                _write_text(cell, value)
        _source_link(data.cell(row_number, 13), fact.source_url)
    data.freeze_panes = "F5"
    data.auto_filter.ref = f"A4:U{max(HEADER_ROW, len(raw_facts) + HEADER_ROW)}"
    for col in range(1, 22):
        data.column_dimensions[get_column_letter(col)].width = {2: 52, 3: 52, 6: 25, 7: 25, 12: 72, 13: 82, 14: 45, 15: 34, 16: 25}.get(col, 20)

    report("Adding filing coverage and missing-history notes…")
    sources = _base_sheet(workbook, "Sources & Coverage", bundle.company.name,
                          f"{len(periods)} selected annual period ends (20 maximum) | {len(raw_facts)} raw facts | "
                          f"{len(bundle.filings)} discovered filings | {len(bundle.coverage)} extraction coverage records", 8)
    _headers(sources, ["Record", "Accession", "Form", "Report end", "Filed", "Status", "Detail", "Source filing"])
    records = []
    forms = {filing.accession: filing.form for filing in bundle.filings}
    for filing in bundle.filings:
        records.append(["Filing", filing.accession, filing.form, filing.report_date, filing.filed,
                        "Discovered", "Metadata only; see extraction coverage records for availability.", filing.source_url])
    for coverage in bundle.coverage:
        records.append(["Coverage", coverage.accession, forms.get(coverage.accession, ""), coverage.report_date,
                        coverage.filed, coverage.status, coverage.detail, coverage.source_url])
    for warning in bundle.warnings:
        records.append(["Warning", "", "", "", "", "Notice", warning, ""])
    for statement, unavailable in unavailable_periods.items():
        for period in unavailable:
            records.append(["Unavailable period", "", "", period, "", "No selected values",
                            f"{statement}: no selected statement values for this period. Its column remains blank in the statement view and Data Sheet, matching the online presentation; values were not estimated. Filing discovery and raw facts remain available here and in Source Data.", ""])
    records.append(["Presentation policy", "", "", "", "", "Latest 20 fiscal periods",
                    "Statement tabs and Data Sheet use the same selected financial facts and a maximum of 20 fiscal periods. Source Data retains all supplied raw observations, including older and unselected facts; those are not additional presented periods.", ""])
    _check_dimensions(len(records) + HEADER_ROW, 8, "Sources & Coverage")
    for row_number, record in enumerate(records, FIRST_DATA_ROW):
        _style_row(sources, row_number, 8)
        for col, value in enumerate(record, 1):
            _write_text(sources.cell(row_number, col), value)
        _source_link(sources.cell(row_number, 8), record[7])
        sources.cell(row_number, 7).alignment = Alignment(wrap_text=True, vertical="top")
    sources.freeze_panes = "C5"
    sources.auto_filter.ref = f"A4:H{max(HEADER_ROW, len(records) + HEADER_ROW)}"
    for col, width in enumerate([15, 26, 12, 16, 16, 24, 88, 85], 1):
        sources.column_dimensions[get_column_letter(col)].width = width

    readme = _base_sheet(workbook, "Read Me", bundle.company.name,
                         "Fishtorian · Historical financial research. Reported facts, transparent calculations, traceable sources.", 2)
    from openpyxl.drawing.image import Image
    from .branding import ASSETS
    logo = Image(str(ASSETS / "fishtorian-workbook.png"))
    logo.width, logo.height = 260, 95
    readme.add_image(logo, "A3")
    readme.row_dimensions[3].height = 76
    _headers(readme, ["Topic", "Explanation"])
    notes = [
        ("Company", bundle.company.name), ("Ticker", bundle.company.ticker), ("SEC CIK", bundle.company.cik),
        ("Generated (UTC)", bundle.generated_at),
        ("Annual history", "Presentation is limited to the latest 20 selected fiscal periods, in chronological order. Statement tabs and Data Sheet retain the same selected columns, including entirely unavailable periods and internal gaps as blanks, matching the online presentation. Sources & Coverage lists unavailable statement periods. Raw observations outside the presentation window may remain in Source Data. Available data does not prove complete lifetime coverage."),
        ("Presentation", "The main statements use annual columns, section headings and bold reported totals. Cash Flow presents reported cash reconciliation totals; Cash Flow Detail retains every selected company cash-flow line. Data Sheet combines Profit & Loss, Balance Sheet and full Cash Flow in a clean annual matrix using the same selected facts, units and signs. No forecasts, valuations or synthetic totals are added."),
        ("Units", "Pure currency amounts display in millions of their own reported currency. Shares, currency per share, ratios and all other units are unscaled. The statement note specifies a single reported currency; mixed currencies and other units are identified in row labels. Each cell comment records its exact reported and display unit. No currency conversion occurs."),
        ("Data Sheet columns", "Column A: reported line item, with units where needed. Columns B through at most U: selected annual fiscal period ends, oldest to newest, repeated for each statement block. Source tags, dates, accessions and links stay in value-cell comments and the separate Source Data audit table."),
        ("Raw values", "Source Data stores every supplied raw fact and every selected statement fact, unscaled. Displayed cells identifies every statement or Data Sheet cell using each selected fact."),
        ("Presentation signs", "Where the filing's visible statement sign is available, statement tabs and Data Sheet use that presentation value. For example, an XBRL payments concept can be positive while its cash-flow line is shown negative. Source Data keeps the underlying Source/XBRL value and optional Statement presentation value separately, both unscaled."),
        ("Cash-flow reconciliation", "Beginning and ending cash can use the same original SEC concept with different instant dates. They remain separate reported rows, identified by cash_begin/cash_end row roles. Annual display period columns identify the verified year being reconciled; Original source instant preserves each actual XBRL context date. These display periods do not turn cash balances into duration facts."),
        ("Missing values", "Blank cells mean unavailable or withheld, never an assumed zero. A reported zero remains numerical zero. A comment explains any fact withheld because its period, unit, tag or statement differs from its row."),
        ("Latest reported history", "Selection aims to use the latest available annual filed representation for a comparable historical period. A later comparative figure is not necessarily a confirmed restatement. Read each cell's selection reason and filing date."),
        ("SEC companyfacts limits", "The SEC companyfacts API contains standardized entity-level XBRL facts, generally from the modern structured reporting era. It is not a complete statement presentation, omits company-specific custom tags and dimensions, and does not establish 50+ years of coverage."),
        ("Older filings", "Older or unsupported filings may have no extractable numerical facts. Sources & Coverage records discovered filings, supported extraction, unsupported formats and reported gaps. Missing years are not estimated."),
        ("Untagged row identifiers", "Identifiers beginning html: are locally generated from an untagged statement label. They are not SEC XBRL tags. Consult the source filing and extraction notes."),
        ("Comparability", "Company labels and SEC tags are preserved where supplied. Only explicitly recognized equivalent concepts share a historical row; their original labels and tags remain in each cell comment and Source Data. Other custom and ambiguous concepts remain separate. Different units remain separate. Fiscal periods may have different ends or lengths; inspect start and end dates."),
        ("EPS and stock splits", "Earnings per share and share counts retain each selected filing’s reported basis. Older years are not uniformly adjusted for subsequent stock splits. Review split dates and source notes before comparing per-share history or using it in projections or ratios."),
        ("Analyst adjustments", "Figures remain as reported in the selected source. No normalization of one-time expenses, valuations, growth forecasts or analyst adjustments is performed."),
        ("Provenance", "Every displayed numeric statement value has a comment with SEC tag, raw unit/value, accession, filed date, period dates, extraction and selection reason. Valid SEC source URLs are clickable."),
        ("Link policy", "Only HTTPS links on sec.gov, www.sec.gov or data.sec.gov using SEC archive/XBRL/submissions paths become hyperlinks. Other source text remains literal and is not executable."),
        ("Excel text safety", "Source and user text is always literal. Only application-generated analytical formulas are executable; their dependency graph is in Calculation Inputs."),
        ("Historical analysis", "Supporting schedules group reported observations. Historical Ratios, Return Bridge and Common Size use documented calculations, not forecasts. Methodology defines applicability and limitations."),
        ("Formula results", "Analytical formulas include cached numerical results for data-only readers. Excel automatic recalculation is enabled; edits require a spreadsheet application that supports these formulas."),
        ("Acceptance status", "Independent live-filing verification across representative companies remains required before institutional acceptance. Passed arithmetic checks are not an audit."),
        ("Withheld inconsistent cells", str(withheld)),
    ]
    for row_number, (topic, explanation) in enumerate(notes, FIRST_DATA_ROW):
        _style_row(readme, row_number, 2)
        _write_text(readme.cell(row_number, 1), topic)
        _write_text(readme.cell(row_number, 2), explanation)
        readme.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[row_number].height = 48
    readme.column_dimensions["A"].width = 29
    readme.column_dimensions["B"].width = 120
    readme.freeze_panes = "B5"

    report("Building historical schedules and auditable calculations…")
    from .analysis import build_analysis
    from .analysis_export import append_analysis, cache_formula_results
    cache = append_analysis(workbook, bundle, analysis if analysis is not None else build_analysis(bundle))
    report("Saving the Excel file…")
    output = BytesIO()
    workbook.save(output)
    report("Excel file is ready.")
    return cache_formula_results(output.getvalue(), workbook, cache)
