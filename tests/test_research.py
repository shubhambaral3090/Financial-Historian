"""Automatic page orchestration is exercised without real provider traffic."""
from dataclasses import replace
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from io import BytesIO
from threading import Event, get_ident

from openpyxl import load_workbook
import pytest

from financial_historian.market import Quote
from financial_historian.models import Company, Fact, StatementBundle, StatementRow
from financial_historian.research import ResearchController


COMPANY = Company("0000000001", "Fictional Research Fixture", "TEST")


def bundle_for(company=COMPANY, value=120):
    bundle = StatementBundle(company)
    facts = [Fact("sample:Revenue", "Reported revenue", "USD", value + year,
                  f"{year}-01-01", f"{year}-12-31", "0000000001-26-000001", "2026-02-01", "10-K",
                  "https://www.sec.gov/Archives/edgar/data/1/annual.htm", "Profit & Loss")
             for year in range(2000, 2026)]
    bundle.raw_facts = facts
    bundle.statements["Profit & Loss"] = [StatementRow("sample:Revenue", "Reported revenue", "USD",
                                                     {fact.end: fact for fact in facts})]
    return bundle


class FixtureService:
    def __init__(self):
        self.resolutions = []
        self.retrievals = []
        self.fail_profile = False
        self.fail_finances = False
        self.value = 120

    def resolve_company(self, company):
        self.resolutions.append(company)
        if self.fail_profile:
            raise ValueError("Private contact must never escape: private-fixture.invalid")
        return COMPANY if not company.cik else company

    def generate(self, company, progress=None):
        self.retrievals.append(company)
        if self.fail_finances:
            raise RuntimeError("Private contact must never escape: private-fixture.invalid")
        progress("Preparing financial statements")
        return bundle_for(company, self.value)


class FixtureMarket:
    def __init__(self):
        self.calls = []
        self.fail = False

    def quote(self, ticker, *, force_refresh=False):
        self.calls.append((ticker, force_refresh))
        if self.fail:
            raise RuntimeError("Raw provider secret or technical details")
        return Quote(ticker, 50.0, 1.0, 2.0, (), "Yahoo Finance", "2026-01-01T00:00:00+00:00",
                     "America/New_York", "latest available close")


class Clock:
    def __init__(self):
        self.now = datetime(2026, 10, 8, 0, 0, tzinfo=timezone.utc)

    def __call__(self):
        return self.now

    def advance(self, seconds):
        self.now += timedelta(seconds=seconds)


def exporter(bundle, progress=None):
    return str([(row.label, tuple((period, fact.value) for period, fact in row.values.items()))
                for rows in bundle.statements.values() for row in rows]).encode()


def test_selection_automatically_builds_exact_twenty_period_display_and_export():
    service, market = FixtureService(), FixtureMarket()
    events = []
    result = ResearchController().load(service, COMPANY, market, on_event=events.append)
    assert result.company == COMPANY
    assert result.quote.price == 50
    assert len(result.bundle.periods) == 20
    assert len(result.bundle.raw_facts) == 26
    assert result.workbook.startswith(b"PK")
    assert {event.kind for event in events} >= {"company", "market", "financial", "workbook", "progress"}
    workbook = load_workbook(BytesIO(result.workbook))
    sheet = workbook["Profit & Loss"]
    exported = {cell.value: cell.column for cell in sheet[4] if cell.column > 1}
    assert set(exported) == set(result.bundle.periods)
    selected = result.bundle.statements["Profit & Loss"][0]
    for period, fact in selected.values.items():
        cell = next(row[exported[period] - 1] for row in sheet.iter_rows(min_row=5)
                    if row[exported[period] - 1].comment and "sample:Revenue" in row[exported[period] - 1].comment.text)
        assert cell.value == fact.value / 1_000_000
        assert fact.source_url in cell.comment.text
    assert result.bundle.generated_at == result.sec_retrieved_at
    assert result.workbook_generated_at
    assert market.calls == [("TEST", True)]


def test_partial_market_result_is_delivered_before_slow_sec_finishes_on_caller_thread():
    service, market = FixtureService(), FixtureMarket()
    market_received = Event()
    original = service.resolve_company
    caller = get_ident()
    thread_ids = []

    def slow_profile(company):
        assert market_received.wait(2), "Market callback was blocked behind SEC profile validation"
        return original(company)

    service.resolve_company = slow_profile

    def receive(event):
        thread_ids.append(get_ident())
        if event.kind == "market":
            market_received.set()

    result = ResearchController(exporter=exporter).load(service, COMPANY, market, on_event=receive)
    assert result.workbook
    assert set(thread_ids) == {caller}


def test_financial_preview_is_delivered_before_slow_market_finishes():
    financial_received = Event()
    market = FixtureMarket()
    original = market.quote

    def slow_quote(ticker, **kwargs):
        assert financial_received.wait(2), "Financial preview was blocked behind Yahoo"
        return original(ticker, **kwargs)

    market.quote = slow_quote
    result = ResearchController(exporter=exporter).load(
        FixtureService(), COMPANY, market,
        on_event=lambda event: financial_received.set() if event.kind == "financial" else None)
    assert result.quote.price == 50
    assert result.workbook


@pytest.mark.parametrize("failure", ["fail_profile", "fail_finances"])
def test_sec_failures_leave_market_available_and_do_not_leak_private_errors(failure):
    service = FixtureService()
    setattr(service, failure, True)
    events = []
    result = ResearchController(exporter=exporter).load(service, COMPANY, FixtureMarket(), on_event=events.append)
    assert result.quote.price == 50
    assert result.bundle is None
    assert result.workbook is None
    assert result.sec_error
    assert "private-fixture" not in str(result) + str(events)
    if failure == "fail_profile":
        assert service.retrievals == []


def test_market_failure_leaves_financial_preview_and_workbook_available():
    market = FixtureMarket()
    market.fail = True
    result = ResearchController(exporter=exporter).load(FixtureService(), COMPANY, market)
    assert result.quote is None
    assert result.market_error
    assert result.bundle and result.workbook
    assert "Raw provider" not in str(result)


def test_provider_unavailable_response_is_nonfatal():
    market = FixtureMarket()
    market.quote = lambda ticker, **kwargs: Quote(ticker, None, None, None, (), "Yahoo Finance", "", "UTC", "unavailable")
    result = ResearchController(exporter=exporter).load(FixtureService(), COMPANY, market)
    assert result.market_error and result.workbook


def test_browser_refresh_rechecks_both_providers_but_widget_cache_does_not_retrieve_sec():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    controller.load(service, COMPANY, market)
    controller.load(service, COMPANY, market, force_refresh=False)
    assert len(service.retrievals) == 1
    controller.load(service, COMPANY, market, force_refresh=True)
    assert len(service.retrievals) == 2
    assert [refresh for _, refresh in market.calls] == [True, False, True]


def test_failed_refresh_retains_last_good_pair_with_original_timestamps_then_recovers():
    clock = Clock()
    controller = ResearchController(exporter=exporter, clock=clock)
    service, market = FixtureService(), FixtureMarket()
    first = controller.load(service, COMPANY, market)
    clock.advance(60)
    service.fail_finances = True
    fallback = controller.load(service, COMPANY, market)
    assert fallback.sec_stale and fallback.sec_error
    assert fallback.workbook == first.workbook
    assert fallback.bundle == first.bundle
    assert fallback.sec_retrieved_at == first.sec_retrieved_at
    assert fallback.workbook_generated_at == first.workbook_generated_at
    assert any(first.sec_retrieved_at in warning for warning in fallback.warnings)
    service.fail_finances = False
    service.value = 900
    recovered = controller.load(service, COMPANY, market)
    assert not recovered.sec_stale
    assert recovered.workbook != first.workbook
    assert recovered.sec_retrieved_at != first.sec_retrieved_at


def test_workbook_failure_never_attaches_old_workbook_to_new_statements():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    previous = controller.load(service, COMPANY, market)
    service.value = 900

    def broken_export(*args, **kwargs):
        raise RuntimeError("Internal technical information")

    controller.exporter = broken_export
    result = controller.load(service, COMPANY, market)
    assert result.bundle != previous.bundle
    assert result.workbook is None
    assert result.workbook_error
    assert not result.workbook_generated_at
    assert "Internal technical" not in str(result)
    # A failed export does not destroy the prior complete successful snapshot.
    service.fail_finances = True
    fallback = controller.load(service, COMPANY, market)
    assert fallback.workbook == previous.workbook
    assert fallback.bundle == previous.bundle


def test_financial_cache_entries_expire_even_when_refresh_fails():
    clock = Clock()
    controller = ResearchController(exporter=exporter, clock=clock, retention_seconds=10)
    service = FixtureService()
    controller.load(service, COMPANY, FixtureMarket())
    clock.advance(10)
    service.fail_finances = True
    assert controller.load(service, COMPANY, FixtureMarket()).bundle is None


def test_financial_cache_is_bounded_by_company_count_and_payload_bytes():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter, max_entries=1)
    controller.load(service, COMPANY, market)
    controller.load(service, replace(COMPANY, cik="0000000002", ticker="TWO"), market)
    service.fail_finances = True
    assert controller.load(service, COMPANY, market).bundle is None
    tiny = ResearchController(exporter=exporter, max_bytes=1)
    service.fail_finances = False
    assert tiny.load(service, COMPANY, market).workbook
    service.fail_finances = True
    assert tiny.load(service, COMPANY, market).bundle is None


def test_cached_snapshot_does_not_share_mutable_bundle_with_browser_session():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    initial = controller.load(service, COMPANY, market)
    initial.bundle.warnings.append("Caller mutated this view")
    initial.bundle.periods.clear()
    service.fail_finances = True
    fallback = controller.load(service, COMPANY, market)
    assert len(fallback.bundle.periods) == 20
    assert "Caller mutated this view" not in fallback.bundle.warnings
    assert all(not isinstance(value, FixtureService) for value in vars(controller).values())


def test_cik_only_selection_starts_market_once_sec_supplies_ticker():
    service, market = FixtureService(), FixtureMarket()
    service.resolve_company = lambda _: COMPANY
    result = ResearchController(exporter=exporter).load(service, replace(COMPANY, ticker=""), market)
    assert market.calls == [("TEST", True)]
    assert result.company == COMPANY
    assert result.quote.price == 50


def test_former_issuer_without_ticker_still_builds_financial_workbook():
    company = replace(COMPANY, ticker="")
    market = FixtureMarket()
    result = ResearchController(exporter=exporter).load(FixtureService(), company, market)
    assert result.workbook
    assert not market.calls
    assert not result.market_error


def test_yahoo_share_class_symbol_is_normalized_only_for_market_request():
    company = replace(COMPANY, ticker="BRK.B")
    market = FixtureMarket()
    result = ResearchController(exporter=exporter).load(FixtureService(), company, market)
    assert result.company.ticker == "BRK.B"
    assert market.calls == [("BRK-B", True)]


@pytest.mark.parametrize("warning", [
    "Structured company-facts data could not be loaded: temporary provider failure",
    "One or more SEC filing documents could not be retrieved; network or provider access prevented a complete refresh.",
])
def test_empty_provider_failure_coverage_does_not_replace_good_statements(warning):
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    previous = controller.load(service, COMPANY, market)
    service.generate = lambda company, progress=None: StatementBundle(company, warnings=[warning])
    result = controller.load(service, COMPANY, market)
    assert result.sec_stale
    assert result.workbook == previous.workbook
    assert result.bundle == previous.bundle


def test_first_attempt_with_no_financial_values_still_delivers_honest_coverage_workbook():
    service = FixtureService()
    service.generate = lambda company, progress=None: StatementBundle(company, warnings=[
        "Structured company-facts data could not be loaded: temporary provider failure"])
    result = ResearchController(exporter=exporter).load(service, COMPANY, FixtureMarket())
    assert result.bundle is not None and result.workbook is not None
    assert not result.sec_stale
    assert not result.bundle.raw_facts


def test_mismatched_ticker_cik_request_cannot_receive_another_cached_company():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    controller.load(service, COMPANY, market)
    service.fail_profile = True
    wrong = controller.load(service, replace(COMPANY, ticker="OTHER"), market)
    assert wrong.bundle is None and wrong.workbook is None
    assert wrong.quote.ticker == "OTHER"


def test_reassigned_ticker_does_not_fall_back_to_previous_issuers_financials():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    controller.load(service, COMPANY, market)
    service.resolve_company = lambda _: replace(COMPANY, cik="0000000002", name="New issuer")
    service.fail_finances = True
    result = controller.load(service, replace(COMPANY, cik=""), market)
    assert result.company.cik == "0000000002"
    assert result.bundle is None and result.workbook is None


def test_partial_network_refresh_is_displayed_without_replacing_complete_last_good_pair():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    complete = controller.load(service, COMPANY, market)
    original_generate = service.generate

    def partial_generate(company, progress=None):
        bundle = bundle_for(company, 999)
        row = bundle.statements["Profit & Loss"][0]
        row.values = {"2025-12-31": row.values["2025-12-31"]}
        bundle.warnings.append("One or more SEC filing documents could not be retrieved; "
                               "network or provider access prevented a complete refresh.")
        return bundle

    service.generate = partial_generate
    partial = controller.load(service, COMPANY, market)
    assert partial.workbook != complete.workbook
    assert len(partial.bundle.statements["Profit & Loss"][0].values) == 1
    assert not partial.sec_stale  # It is newly retrieved partial data, clearly warned.
    service.generate = original_generate
    service.fail_finances = True
    failed = controller.load(service, COMPANY, market)
    assert failed.workbook == complete.workbook
    assert failed.bundle == complete.bundle
    assert failed.sec_stale


def test_first_network_partial_is_viewable_but_not_retained_as_last_good():
    service, market = FixtureService(), FixtureMarket()
    controller = ResearchController(exporter=exporter)
    original_generate = service.generate

    def partial_generate(company, progress=None):
        bundle = bundle_for(company)
        bundle.warnings.append("Structured company-facts data could not be loaded: provider unavailable")
        return bundle

    service.generate = partial_generate
    assert controller.load(service, COMPANY, market).workbook
    service.generate = original_generate
    service.fail_finances = True
    assert controller.load(service, COMPANY, market).workbook is None


def test_late_older_request_cannot_overwrite_newer_successful_snapshot():
    controller = ResearchController(exporter=exporter)
    old_service, new_service = FixtureService(), FixtureService()
    old_service.value = 100
    new_service.value = 900
    older_started, release_older = Event(), Event()
    original_generate = old_service.generate

    def slow_generate(company, progress=None):
        older_started.set()
        assert release_older.wait(3)
        return original_generate(company, progress)

    old_service.generate = slow_generate
    with ThreadPoolExecutor(max_workers=1) as executor:
        older_future = executor.submit(controller.load, old_service, COMPANY, FixtureMarket())
        assert older_started.wait(2)
        try:
            newer = controller.load(new_service, COMPANY, FixtureMarket())
        finally:
            release_older.set()
        older = older_future.result(timeout=3)
    assert older.workbook != newer.workbook
    cached = controller.load(new_service, COMPANY, FixtureMarket(), force_refresh=False)
    assert cached.workbook == newer.workbook
    assert cached.bundle == newer.bundle


def test_canonical_validated_presentation_is_shared_by_callback_result_and_excel():
    service = FixtureService()

    def inconsistent(company, progress=None):
        bundle = bundle_for(company)
        row = bundle.statements["Profit & Loss"][0]
        row.values["2025-12-31"] = replace(row.values["2025-12-31"], unit="shares")
        return bundle

    service.generate = inconsistent
    events = []
    result = ResearchController().load(service, COMPANY, FixtureMarket(), on_event=events.append)
    row = result.bundle.statements["Profit & Loss"][0]
    assert "2025-12-31" not in row.values
    preview = next(event.bundle for event in events if event.kind == "financial")
    assert preview == result.bundle
    assert any("Withheld from statement presentation" in warning for warning in result.bundle.warnings)
    assert any(fact.unit == "shares" for fact in result.bundle.raw_facts)
    workbook = load_workbook(BytesIO(result.workbook))
    sheet = workbook["Profit & Loss"]
    column = next(cell.column for cell in sheet[4] if cell.value == "2025-12-31")
    revenue_row = next(cells for cells in sheet.iter_rows(min_row=5) if cells[0].value == "Reported revenue")
    assert revenue_row[column - 1].value is None
