"""A complete mocked SEC lookup -> filing parse -> analyst workbook flow."""
from io import BytesIO

from openpyxl import load_workbook

from financial_historian.export import build_workbook
from financial_historian.models import STATEMENTS
from financial_historian.service import FinancialHistorian


def document(year, revenue, comparative=None):
    context = lambda y: f'''<xbrli:context id="y{y}"><xbrli:entity><xbrli:identifier scheme="http://www.sec.gov/CIK">123</xbrli:identifier></xbrli:entity><xbrli:period><xbrli:startDate>{y}-01-01</xbrli:startDate><xbrli:endDate>{y}-12-31</xbrli:endDate></xbrli:period></xbrli:context>'''
    instant = f'''<xbrli:context id="instant"><xbrli:entity><xbrli:identifier scheme="http://www.sec.gov/CIK">123</xbrli:identifier></xbrli:entity><xbrli:period><xbrli:instant>{year}-12-31</xbrli:instant></xbrli:period></xbrli:context>'''
    amount = lambda tag, value, cid: f'<ix:nonFraction name="{tag}" unitRef="USD" contextRef="{cid}" scale="6" format="ixt:num-dot-decimal">{value}</ix:nonFraction>'
    compare = '' if comparative is None else f'<td>{amount("sample:ReportedRevenue", comparative, "y2024")}</td>'
    return f'''<html xmlns:ix="http://www.xbrl.org/2013/inlineXBRL" xmlns:xbrli="http://www.xbrl.org/2003/instance"><body>
    <ix:header>{context(year)}{context(2024) if year != 2024 else ''}{instant}<xbrli:unit id="USD"><xbrli:measure>iso4217:USD</xbrli:measure></xbrli:unit></ix:header>
    <h2>CONSOLIDATED STATEMENTS OF OPERATIONS</h2><table><tr><td>Reported revenue</td><td>{amount('sample:ReportedRevenue', revenue, f'y{year}')}</td>{compare}</tr></table>
    <h2>CONSOLIDATED BALANCE SHEETS</h2><table><tr><td>Total assets</td><td>{amount('us-gaap:Assets', 250, 'instant')}</td></tr></table>
    <h2>CONSOLIDATED STATEMENTS OF CASH FLOWS</h2><table><tr><td>Capital expenditure</td><td>({amount('us-gaap:PaymentsToAcquirePropertyPlantAndEquipment', 20, f'y{year}')})</td></tr></table>
    </body></html>'''.encode()


def test_entire_workflow_keeps_revisions_custom_tags_and_printed_cash_sign(monkeypatch, tmp_path):
    annuals = {
        'accessionNumber': ['0000000123-25-000001', '0000000123-26-000001'],
        'form': ['10-K', '10-K'], 'filingDate': ['2025-02-01', '2026-02-01'],
        'reportDate': ['2024-12-31', '2025-12-31'], 'primaryDocument': ['2024.htm', '2025.htm'],
    }
    class FixtureClient:
        def __init__(self, *args, **kwargs):
            pass
        def get_json(self, url, **kwargs):
            if 'company_tickers' in url:
                return {'0': {'cik_str': 123, 'title': 'Fictional Workflow Test Co', 'ticker': 'TEST'}}
            if 'companyfacts' in url:
                return {'cik': 123, 'facts': {}}
            return {'filings': {'recent': annuals, 'files': []}}
        def get_bytes(self, url, **kwargs):
            return document(2024, 100) if url.endswith('2024.htm') else document(2025, 120, 90)
    monkeypatch.setattr('financial_historian.service.SecClient', FixtureClient)
    service = FinancialHistorian('Mocked test identity', tmp_path)
    company = service.search_companies('TEST')[0]
    messages = []
    bundle = service.generate(company, progress=messages.append)
    assert set(bundle.statements) == set(STATEMENTS)
    assert len(bundle.coverage) == 2
    # One row per statement is a partial parse, not a complete statement.
    assert all(row.status == 'partial' for row in bundle.coverage)
    assert all('Core rows not verified' in row.detail for row in bundle.coverage)
    workbook = load_workbook(BytesIO(build_workbook(bundle)))
    revenue = next(row for row in bundle.statements['Profit & Loss'] if row.tag == 'sample:ReportedRevenue')
    assert revenue.values['2024-12-31'].value == 90_000_000
    assert revenue.values['2024-12-31'].filed == '2026-02-01'
    assert any('earlier' in warning for warning in bundle.warnings)
    def sourced_cell(sheet, tag, period):
        column = next(cell.column for cell in sheet[4] if cell.value == period)
        return next(row[column - 1] for row in sheet.iter_rows(min_row=5)
                    if row[column - 1].comment and tag in row[column - 1].comment.text)

    assert sourced_cell(workbook['Profit & Loss'], 'sample:ReportedRevenue', '2024-12-31').value == 90
    capex = sourced_cell(workbook['Cash Flow Detail'], 'us-gaap:PaymentsToAcquirePropertyPlantAndEquipment', '2025-12-31')
    assert capex.value == -20
    assert any(fact.value == 20_000_000 and fact.presentation_value == -20_000_000 for fact in bundle.raw_facts)
    assert capex.comment and '2026-02-01' in capex.comment.text
    assert messages and 'all available' in messages[0]
