from dataclasses import replace
from decimal import Decimal
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from financial_historian.analysis import build_analysis, Calculator
from financial_historian.export import build_workbook
from financial_historian.models import Fact, StatementRow
from financial_historian.operating_kpis import KpiMapping, load_kpi_mappings
from financial_historian.disclosures import parse_disclosures
from test_analysis import historical_fixture, cell, mutate


def metric_fixture():
    bundle = historical_fixture()
    specs = [
        ("Profit & Loss", "IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest", "USD", [75,87.5,100,150]),
        ("Profit & Loss", "IncomeTaxExpenseBenefit", "USD", [15,17.5,20,30]),
        ("Balance Sheet", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest", "USD", [410,460,510,620]),
        ("Cash Flow", "DepreciationDepletionAndAmortization", "USD", [20,22,25,30]),
        ("Cash Flow", "PaymentsOfDividendsCommonStock", "USD", [10,12,15,20]),
        ("Cash Flow", "PaymentsForRepurchaseOfCommonStock", "USD", [20,25,30,40]),
        ("Profit & Loss", "CommonStockDividendsPerShareDeclared", "USD/shares", [.1,.12,.15,.25]),
        ("Profit & Loss", "EarningsPerShareBasic", "USD/shares", [.58,.68,.78,1.475]),
        ("Profit & Loss", "EarningsPerShareDiluted", "USD/shares", [.58,.68,.78,1.475]),
        ("Profit & Loss", "WeightedAverageNumberOfDilutedSharesOutstanding", "shares", [100,100,100,80]),
    ]
    for statement,tag,unit,values in specs:
        row = StatementRow("us-gaap:"+tag,tag,unit)
        for year,value in zip(range(2021,2025),values):
            end = f"{year}-12-31"
            source = bundle.statements["Profit & Loss"][0].values[end]
            fact = replace(source,tag=row.tag,label=tag,unit=unit,
                           value=value if unit=="USD/shares" else value*1000000,
                           start=None if statement=="Balance Sheet" else f"{year}-01-01",
                           statement=statement,decimals="3" if unit=="USD/shares" else "-3")
            row.values[end] = fact
            bundle.raw_facts.append(fact)
        bundle.statements[statement].append(row)
    # Independent current-filing comparative observations; no split adjustment is
    # inferred by the implementation from different filings.
    for fact in list(bundle.raw_facts):
        if fact.end == "2023-12-31":
            current_source=bundle.statements["Profit & Loss"][0].values["2024-12-31"]
            bundle.raw_facts.append(replace(fact, accession=current_source.accession,
                                           filed=current_source.filed, source_url=current_source.source_url))
    sample = bundle.statements["Profit & Loss"][0].values["2024-12-31"]
    for year,value in [(2023,400000000),(2024,600000000)]:
        bundle.disclosures.append(replace(sample,start=f"{year}-01-01",end=f"{year}-12-31",value=value,
            label="Retail revenue",statement="Disclosure",section="Revenue & Segments",
            dimensions=(("demo:BusinessSegmentsAxis","demo:RetailMember"),),context_id=f"retail{year}"))
    for year,value in [(2023,10),(2024,12)]:
        bundle.disclosures.append(replace(sample,tag="demo:ActiveAccounts",label="Active accounts",unit="pure",start=None,
            end=f"{year}-12-31",value=value,decimals="0",statement="Disclosure",section="Operating KPIs",
            row_role="kpi:demo:ActiveAccounts",selection_reason="Reviewed KPI: active accounts at fiscal end, not unique people."))
    return bundle


@pytest.mark.parametrize("key,expected",[
    ("ebitda",192000000), ("ebitda_margin",.16), ("debt_ebitda",.9375),
    ("ebitda_interest",16), ("debt_capital",.225), ("equity_assets",620/1300),
    ("effective_tax",.2), ("nopat_proxy",144000000),
    ("invested_capital_proxy",740000000), ("roic_proxy",144/675),
    ("parent_earnings_growth",118/78-1), ("common_earnings_growth",118/78-1),
    ("diluted_eps_growth",1.475/.78-1), ("diluted_shares_growth",-.2),
    ("eps_earnings_gap",1.475/.78-118/78),
    ("cash_dividend_payout",20/118), ("declared_dps_payout",.25/1.475),
    ("repurchase_earnings",40/118), ("common_cash_returned",60000000),
    ("gross_margin_change",500/1200-.4), ("operating_margin_change",.05),
    ("net_margin_change",118/1200-.078),
])
def test_independent_hand_calculated_metrics(key,expected):
    item = cell(build_analysis(metric_fixture()),key)
    assert item.value is not None, item.reason
    assert float(item.value) == pytest.approx(expected)
    assert item.sources


@pytest.mark.parametrize("tag,change,key,reason",[
    ("IncomeTaxExpenseBenefit",{"value":31000000},"ebitda","reconcile"),
    ("InterestExpense",{"value":-1},"ebitda","Negative"),
    ("DepreciationDepletionAndAmortization",{"value":-1},"ebitda","Negative"),
    ("DepreciationDepletionAndAmortization",{"accession":"different"},"ebitda","vintages"),
    ("IncomeTaxExpenseBenefit",{"unit":"EUR"},"ebitda","currencies"),
    ("IncomeTaxExpenseBenefit",{"start":"2024-01-02"},"ebitda","starts"),
    ("InterestExpense",{"value":0},"ebitda_interest","Denominator"),
    ("DebtAndCapitalLeaseObligations",{"value":-1},"debt_capital","Negative"),
    ("StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest",{"value":-1},"debt_capital","Negative"),
    ("CashAndCashEquivalentsAtCarryingValue",{"value":900000000},"roic_proxy","positive"),
    ("NetIncomeLossAvailableToCommonStockholdersBasic",{"value":0},"cash_dividend_payout","Denominator"),
    ("CommonStockDividendsPerShareDeclared",{"unit":"EUR/shares"},"declared_dps_payout","identical"),
    ("EarningsPerShareBasic",{"value":.0001},"declared_dps_payout","precision"),
    ("IncomeTaxExpenseBenefit",{"unit":"shares"},"ebitda","unit"),
    ("WeightedAverageNumberOfDilutedSharesOutstanding",{"value":-1},"diluted_shares_growth","positive"),
    ("Revenues",{"unit":"pure"},"ebitda_margin","units"),
])
def test_incompatible_or_invalid_inputs_remain_explained_blanks(tag,change,key,reason):
    bundle = metric_fixture()
    mutate(bundle,tag,**change)
    if "accession" in change:
        bundle.raw_facts = [f for f in bundle.raw_facts if not (f.tag.endswith(":"+tag) and f.end=="2024-12-31")]
    item = cell(build_analysis(bundle),key)
    assert item.value is None
    assert reason.lower() in item.reason.lower()


def test_ebitda_does_not_substitute_pp_and_e_only_depreciation_or_add_sbc():
    bundle = metric_fixture()
    bundle.statements["Cash Flow"] = [r for r in bundle.statements["Cash Flow"] if not r.tag.endswith(":DepreciationDepletionAndAmortization")]
    assert cell(build_analysis(bundle),"ebitda").value is None


def test_effective_tax_out_of_bounds_and_negative_ebitda():
    bundle = metric_fixture()
    mutate(bundle,"ProfitLoss",value=-200000000)
    mutate(bundle,"IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest",value=-170000000)
    result = build_analysis(bundle)
    assert cell(result,"ebitda").value == -128000000
    assert cell(result,"debt_ebitda").value is None
    assert cell(result,"roic_proxy").value is None
    mutate(bundle,"IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest",value=10000000)
    mutate(bundle,"IncomeTaxExpenseBenefit",value=210000000)
    assert "0–100%" in cell(build_analysis(bundle),"effective_tax").reason


def test_negative_equity_is_retained_in_equity_assets_not_capital_ratios():
    bundle=metric_fixture()
    mutate(bundle,"StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest",value=-10000000)
    result=build_analysis(bundle)
    assert cell(result,"equity_assets").value < 0
    assert cell(result,"debt_capital").value is None
    assert cell(result,"roic_proxy").value is None


def test_unknown_or_financial_sector_withholds_industrial_extensions():
    for sic in ("", "6021"):
        bundle=metric_fixture()
        bundle.company=replace(bundle.company,sic=sic)
        result=build_analysis(bundle)
        for key in ("ebitda","debt_ebitda","roic_proxy"):
            assert "financial" in cell(result,key).reason
        assert cell(result,"cash_dividend_payout").value is not None


def test_eps_and_share_growth_require_same_filing_not_different_split_bases():
    bundle=metric_fixture()
    bundle.raw_facts=[f for f in bundle.raw_facts if not (f.end=="2023-12-31" and f.accession=="0000000001-24-000001")]
    result=build_analysis(bundle)
    for key in ("diluted_eps_growth","diluted_shares_growth","eps_earnings_gap"):
        assert cell(result,key).value is None
        assert "same-filing" in cell(result,key).reason


def test_margin_change_rejects_changed_revenue_scope():
    bundle=metric_fixture()
    row=bundle.statements["Profit & Loss"][0]
    previous=row.values.pop("2023-12-31")
    new=replace(previous,tag="us-gaap:SalesRevenueNet")
    bundle.statements["Profit & Loss"].append(StatementRow(new.tag,"Net sales","USD",{new.end:new}))
    item=cell(build_analysis(bundle),"operating_margin_change")
    assert item.value is None and "basis changed" in item.reason


def test_roic_uses_sourced_opening_capital_outside_display_window():
    bundle=metric_fixture()
    bundle.periods=["2024-12-31"]
    for rows in bundle.statements.values():
        for row in rows:
            row.values={p:f for p,f in row.values.items() if p=="2024-12-31"}
    item=cell(build_analysis(bundle),"roic_proxy")
    assert float(item.value)==pytest.approx(144/675)
    assert any(f.end=="2023-12-31" for f in item.sources)


def test_missing_common_dividends_are_not_substituted_with_total_dividends():
    bundle=metric_fixture()
    row=next(r for r in bundle.statements["Cash Flow"] if r.tag.endswith(":PaymentsOfDividendsCommonStock"))
    row.tag="us-gaap:PaymentsOfDividends"
    row.values={p:replace(f,tag=row.tag) for p,f in row.values.items()}
    assert cell(build_analysis(bundle),"cash_dividend_payout").value is None


def test_cash_distribution_ratios_respect_printed_outflow_signs():
    bundle=metric_fixture()
    mutate(bundle,"PaymentsOfDividendsCommonStock",presentation_value=-20000000)
    mutate(bundle,"PaymentsForRepurchaseOfCommonStock",presentation_value=-40000000)
    result=build_analysis(bundle)
    assert cell(result,"common_cash_returned").value==60000000
    assert float(cell(result,"cash_dividend_payout").value)==pytest.approx(20/118)


def test_business_member_rename_is_not_stitched_into_growth():
    bundle=metric_fixture()
    bundle.disclosures=[replace(f,dimensions=(("demo:BusinessSegmentsAxis","demo:RenamedMember"),))
                        if f.label=="Retail revenue" and f.end=="2024-12-31" else f for f in bundle.disclosures]
    rows=[r for r in build_analysis(bundle).rows if "Retail revenue" in r.definition.label and r.definition.key.endswith(":growth")]
    assert rows and all(r.cells["2024-12-31"].value is None for r in rows)


def test_per_share_and_currency_sources_cannot_hide_cross_currency_growth_gap():
    bundle=metric_fixture()
    row=next(r for r in bundle.statements["Profit & Loss"] if r.tag.endswith(":EarningsPerShareDiluted"))
    row.unit="EUR/shares"
    row.values={p:replace(f,unit="EUR/shares") for p,f in row.values.items()}
    bundle.raw_facts=[replace(f,unit="EUR/shares") if f.tag==row.tag else f for f in bundle.raw_facts]
    item=cell(build_analysis(bundle),"eps_earnings_gap")
    assert item.value is None and "currencies" in item.reason


def test_segment_mix_growth_and_kpi_counts_have_distinct_semantics():
    result=build_analysis(metric_fixture())
    rows=[r for r in result.rows if r.definition.group=="Business drivers"]
    assert len(rows)==3
    assert [float(r.cells["2024-12-31"].value) for r in rows]==[.5,.5,.2]
    assert all(r.cells["2024-12-31"].sources for r in rows)
    assert "not unique people" in rows[-1].definition.purpose


def test_segment_mix_rejects_nonmonetary_consolidated_revenue_unit():
    bundle=metric_fixture()
    mutate(bundle,"Revenues",unit="pure")
    result=build_analysis(bundle)
    row=next(r for r in result.rows if r.definition.key.endswith(":mix"))
    assert row.cells["2024-12-31"].value is None
    assert "units" in row.cells["2024-12-31"].reason


@pytest.mark.parametrize("changes",[
    {"value":1300000000}, {"value":-1}, {"accession":"different"},
    {"start":"2024-01-02"}, {"unit":"EUR"},
])
def test_segment_mix_rejects_eliminations_scope_currency_and_vintage(changes):
    bundle=metric_fixture()
    bundle.disclosures=[replace(f,**changes) if f.label=="Retail revenue" and f.end=="2024-12-31" else f for f in bundle.disclosures]
    rows=[r for r in build_analysis(bundle).rows if r.definition.key.endswith(":mix")]
    assert rows and all(r.cells["2024-12-31"].value is None for r in rows)


def test_all_new_export_values_match_shared_graph_with_formulas_and_provenance():
    bundle=metric_fixture()
    analysis=build_analysis(bundle)
    payload=build_workbook(bundle)
    cached=load_workbook(BytesIO(payload),data_only=True)
    formula=load_workbook(BytesIO(payload))
    assert {"EBITDA & Capital","Operating KPIs","Business Drivers"} <= set(cached.sheetnames)
    for metric in analysis.rows:
        item=metric.cells["2024-12-31"]
        if item.value is None or metric.definition.group=="Return bridge":
            continue
        row=next(r for r in cached["Historical Ratios"] if str(r[0].value).startswith(metric.definition.label+" ["))
        expected=float(item.value)*(100 if item.unit=="pp" else 1)/(1000000 if item.unit=="USD" else 1)
        assert row[-1].value==pytest.approx(expected)
        original=formula["Historical Ratios"][row[-1].coordinate]
        assert original.data_type=="f" and "Calculation Inputs!" in original.comment.text
    assert all(sheet.max_column<=21 for sheet in cached if sheet.title in {"Historical Ratios","Business Drivers"})


def mapping():
    return KpiMapping("0000000001","demo:ActiveAccounts","Active accounts","pure","instant",
                      "Active account count at fiscal year end, not unique people.",
                      "https://www.sec.gov/Archives/edgar/data/1/annual.htm","2026-10-08")


def test_reviewed_registry_exact_fields_issuer_and_evidence(tmp_path):
    from dataclasses import asdict
    path=tmp_path/"kpis.json"
    path.write_text(json.dumps({"schema_version":1,"mappings":[asdict(mapping())]}))
    assert load_kpi_mappings("1",path)==(mapping(),)
    assert load_kpi_mappings("2",path)==()
    bad=asdict(mapping())
    bad["source_url"]="https://www.sec.gov/Archives/edgar/data/2/annual.htm"
    path.write_text(json.dumps({"schema_version":1,"mappings":[bad]}))
    with pytest.raises(ValueError,match="issuer"):
        load_kpi_mappings("1",path)


@pytest.mark.parametrize("payload",[
    [], {"schema_version":2,"mappings":[]},
    {"schema_version":1,"mappings":[{"tag":"guess"}]},
])
def test_malformed_kpi_registry_fails_closed(tmp_path,payload):
    path=tmp_path/"kpis.json"
    path.write_text(json.dumps(payload))
    with pytest.raises(ValueError):
        load_kpi_mappings("1",path)


def test_kpi_extraction_requires_exact_approved_tag_unit_and_period_type():
    from financial_historian.models import Filing
    body=b'''<html><xbrli:context id="year"><xbrli:period><xbrli:instant>2024-12-31</xbrli:instant></xbrli:period></xbrli:context>
      <xbrli:unit id="count"><xbrli:measure>xbrli:pure</xbrli:measure></xbrli:unit>
      <table><tr><td>Active accounts</td><td><ix:nonFraction name="demo:ActiveAccounts" contextRef="year" unitRef="count" decimals="0">12</ix:nonFraction></td></tr></table></html>'''
    filing=Filing("a","10-K","2025-02-01","2024-12-31","annual.htm",mapping().source_url)
    assert parse_disclosures(body,filing)==[]
    result=parse_disclosures(body,filing,(mapping(),))
    assert len(result)==1 and result[0].value==12
    assert result[0].section=="Operating KPIs" and result[0].row_role=="kpi:demo:ActiveAccounts"
    assert "not unique people" in result[0].selection_reason
    for wrong in (replace(mapping(),unit="shares"),replace(mapping(),period_type="annual"),replace(mapping(),tag="demo:Subscriptions")):
        assert parse_disclosures(body,filing,(wrong,))==[]
