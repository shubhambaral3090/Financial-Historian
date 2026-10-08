"""The browser and Excel receive the same validated financial cell assignments."""
from dataclasses import replace
from io import BytesIO

from openpyxl import load_workbook
import pytest

from financial_historian.export import build_workbook
from financial_historian.models import Company, Fact, StatementBundle, StatementRow
from financial_historian.presentation import fact_rejection, prepare_presentation


def reported_fact(**changes):
    original = Fact("sample:Revenue", "Reported revenue", "USD", 1200000,
                    "2025-01-01", "2025-12-31", "0000000001-26-000001", "2026-02-01", "10-K",
                    "https://www.sec.gov/Archives/edgar/data/1/annual.htm", "Profit & Loss")
    return replace(original, **changes)


def bundle_for(fact, *, accepted_tags=None):
    bundle = StatementBundle(Company("0000000001", "Presentation Fixture", "TEST"),
                             periods=["2025-12-31"])
    bundle.statements["Profit & Loss"] = [StatementRow("sample:Revenue", "Reported revenue", "USD",
        {"2025-12-31": fact}, accepted_tags=accepted_tags or set(), section="Operating results", is_total=True)]
    return bundle


@pytest.mark.parametrize(("changes", "reason"), [
    ({"end": "2024-12-31"}, "differs from column period"),
    ({"unit": "EUR"}, "units were not mixed"),
    ({"tag": "sample:DifferentScopeRevenue"}, "differs from its row tag"),
    ({"statement": "Balance Sheet"}, "differs from its worksheet"),
])
def test_invalid_assignment_is_blank_in_both_views_and_original_source_is_retained(changes, reason):
    fact = reported_fact(**changes)
    original = bundle_for(fact)
    prepared = prepare_presentation(original)
    assert prepared.statements["Profit & Loss"][0].values == {}
    assert original.statements["Profit & Loss"][0].values == {"2025-12-31": fact}
    assert original.raw_facts == [] and original.warnings == []
    assert prepared.periods == ["2025-12-31"]
    assert prepared.raw_facts == [fact]
    assert reason in prepared.warnings[0]
    workbook = load_workbook(BytesIO(build_workbook(prepared)))
    # Same period and blank selected value, not a relabeled or shifted number.
    assert workbook["Profit & Loss"]["B4"].value == "2025-12-31"
    assert not any(cell.data_type == "n" and cell.value is not None
                   for row in workbook["Profit & Loss"].iter_rows(min_row=5) for cell in row)
    assert workbook["Data Sheet"]["B5"].value == "2025-12-31"
    assert not any(cell.data_type == "n" and cell.value is not None
                   for row in workbook["Data Sheet"].iter_rows(min_row=6) for cell in row)
    sources = workbook["Source Data"]
    assert sources["B5"].value == fact.tag
    assert sources["D5"].value == fact.unit
    assert sources["E5"].value == fact.value
    assert sources["G5"].value == fact.end
    assert sources["N5"].value is None  # It was not displayed in any statement cell.
    coverage = [str(cell.value) for row in workbook["Sources & Coverage"] for cell in row]
    assert prepared.warnings[0] in coverage


def test_valid_fact_keeps_reported_label_source_sign_scope_and_units():
    fact = reported_fact(presentation_value=-1200000)
    prepared = prepare_presentation(bundle_for(fact))
    row = prepared.statements["Profit & Loss"][0]
    assert row.values["2025-12-31"] is fact
    assert row.label == fact.label
    assert row.unit == "USD" and row.is_total
    assert row.section == "Operating results"
    assert prepared.warnings == []
    workbook = load_workbook(BytesIO(build_workbook(prepared)))
    displayed = [cell for record in workbook["Profit & Loss"].iter_rows(min_row=5)
                 for cell in record if cell.comment and fact.tag in cell.comment.text]
    assert len(displayed) == 1 and displayed[0].value == -1.2
    assert fact.source_url in displayed[0].comment.text
    assert workbook["Source Data"]["E5"].value == 1200000
    assert workbook["Source Data"]["O5"].value == -1200000


def test_only_explicitly_accepted_equivalent_tags_can_share_a_row():
    fact = reported_fact(tag="sample:LegacyEquivalentRevenue")
    original = bundle_for(fact, accepted_tags={fact.tag})
    prepared = prepare_presentation(original)
    assert prepared.statements["Profit & Loss"][0].values == {"2025-12-31": fact}
    assert fact_rejection(prepared.statements["Profit & Loss"][0], fact.end, fact, fact.statement) is None
    assert not prepared.warnings
    prepared.statements["Profit & Loss"][0].accepted_tags.clear()
    assert original.statements["Profit & Loss"][0].accepted_tags == {fact.tag}


def test_prepare_is_idempotent_without_duplicate_provenance_or_warnings():
    invalid = reported_fact(unit="EUR")
    original = bundle_for(invalid)
    original.raw_facts = [invalid]
    first = prepare_presentation(original)
    second = prepare_presentation(first)
    assert second == first
    assert len(second.raw_facts) == 1
    assert len(second.warnings) == 1


def test_valid_zero_and_empty_periods_are_preserved_not_filled():
    fact = reported_fact(value=0)
    original = bundle_for(fact)
    original.periods.insert(0, "2024-12-31")
    prepared = prepare_presentation(original)
    assert prepared.periods == original.periods
    assert prepared.statements["Profit & Loss"][0].values == {"2025-12-31": fact}
    workbook = load_workbook(BytesIO(build_workbook(prepared)))
    displayed = [cell for record in workbook["Profit & Loss"].iter_rows(min_row=5)
                 for cell in record if cell.comment and fact.tag in cell.comment.text]
    assert len(displayed) == 1 and displayed[0].value == 0
    assert displayed[0].column == 3


def test_existing_raw_and_selected_source_observations_all_remain_available():
    selected = reported_fact(unit="EUR")
    old = reported_fact(end="1999-12-31", start="1999-01-01", value=123)
    original = bundle_for(selected)
    original.raw_facts = [old]
    prepared = prepare_presentation(original)
    assert prepared.raw_facts == [old, selected]
    assert prepared.statements["Profit & Loss"][0].values == {}
    assert original.raw_facts == [old]
