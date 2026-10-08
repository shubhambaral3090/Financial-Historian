"""Real Chromium, fixture-backed smoke test; not collected by default pytest.

Install optional playwright==1.56.0, then run:
  .venv/bin/python tests/browser_smoke.py --chromium /usr/bin/chromium

No live SEC/Yahoo calls. No contact-input screenshots, logs, or recordings.
Outputs contain only the fictional public research page and downloaded XLSX.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
from urllib.request import urlopen
from zipfile import ZipFile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect


def run(chromium: str, output: Path) -> dict:
    root = Path(__file__).resolve().parents[1]
    with socket.socket() as reservation:
        reservation.bind(("127.0.0.1", 0))
        port = reservation.getsockname()[1]
    base = f"http://127.0.0.1:{port}"
    environment = dict(os.environ)
    environment.pop("SEC_USER_AGENT", None)
    server = subprocess.Popen(
        [sys.executable, "-m", "streamlit", "run", "tests/browser_fixture_app.py",
         "--server.headless=true", "--browser.gatherUsageStats=false",
         "--server.address=127.0.0.1", f"--server.port={port}"],
        cwd=root, env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    report = {"provider_mode": "offline fictional fixture", "viewports": []}
    try:
        for _ in range(100):
            if server.poll() is not None:
                raise AssertionError("Fixture browser server stopped before startup")
            try:
                with urlopen(base + "/_stcore/health", timeout=1) as response:
                    if response.status == 200:
                        break
            except OSError:
                time.sleep(0.1)
        else:
            raise AssertionError("Fixture browser server did not become ready")
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path=chromium, headless=True,
                                                  args=["--no-sandbox", "--disable-dev-shm-usage"])
            for name, width, height in (("desktop", 1440, 1100), ("mobile", 390, 844)):
                context = browser.new_context(viewport={"width": width, "height": height}, accept_downloads=True)
                external = []
                def restrict(route):
                    if route.request.url.startswith(base + "/") or route.request.url.startswith(("data:", "blob:")):
                        route.continue_()
                    else:
                        external.append(route.request.url.split("?", 1)[0])
                        route.abort()
                context.route("**/*", restrict)
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(base)
                dialog = page.get_by_role("dialog")
                expect(dialog).to_be_visible(timeout=20000)
                bounds = dialog.bounding_box()
                assert bounds and bounds["x"] >= 0 and bounds["x"] + bounds["width"] <= width + 1
                assert abs(bounds["x"] + bounds["width"] / 2 - width / 2) < 30
                assert not page.get_by_role("button", name="Search companies", exact=True).count()
                page.get_by_label("Your name or organization", exact=True).fill("Offline Fixture")
                page.get_by_label("Contact email", exact=True).fill("offline@example.invalid")
                page.get_by_role("button", name="Continue to Financial Historian", exact=True).click()
                expect(dialog).to_be_hidden(timeout=15000)
                search = page.get_by_label("Company name, ticker, or CIK", exact=True)
                expect(search).to_be_visible(timeout=15000)
                assert not page.get_by_label("Contact email", exact=True).count()
                search.fill("FIX")
                page.get_by_role("button", name="Search companies", exact=True).click()
                page.get_by_role("button", name="Fictional UI Fixture · FIX", exact=True).click()
                download_button = page.get_by_role("button", name="Download Excel workbook", exact=True)
                expect(download_button).to_be_visible(timeout=20000)
                page.get_by_role("tab", name="Cash flow", exact=True).click()
                expect(page.locator(".statement-preview").filter(has_text="Capital expenditures")).to_be_visible()
                expect(page.locator(".statement-preview").filter(has_text="Capital expenditures")).to_contain_text("(1.2)")
                for label in ("Prepare Excel download", "Refresh market data", "Refresh this company", "Refresh company universe"):
                    assert not page.get_by_role("button", name=label, exact=True).count()
                body = page.locator("body").inner_text()
                assert "offline@example.invalid" not in body and "Offline Fixture" not in body
                assert "102.00 USD" in body
                expect(page.locator('[data-testid="stVegaLiteChart"]')).not_to_have_count(0)
                assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1")
                assert "offline" not in page.url
                assert page.evaluate("Object.keys(localStorage).every(key => !key.includes('sec-contact'))")
                # Streamlit may set its standard XSRF cookie; the identity must
                # not be placed in that cookie or any new contact cookie.
                assert all("sec-contact" not in item["name"] and "offline" not in item["value"].lower()
                           for item in context.cookies())
                with page.expect_download() as item:
                    download_button.click()
                downloaded = item.value
                assert downloaded.suggested_filename == "FIX_Financial_History.xlsx"
                path = output / f"{name}-fixture.xlsx"
                downloaded.save_as(path)
                content = path.read_bytes()
                workbook = load_workbook(BytesIO(content))
                assert {"Data Sheet", "Source Data", "Sources & Coverage"} <= set(workbook.sheetnames)
                values = [cell.value for row in workbook["Cash Flow Detail"] for cell in row]
                assert values.count(-1.2) == 20
                assert "2004-12-31" not in values and "2005-12-31" in values
                with ZipFile(BytesIO(content)) as archive:
                    assert not any(b"offline@example.invalid" in archive.read(member) for member in archive.namelist())
                # Browser-side sessionStorage must survive an actual document
                # reload, while Python session_state and provider checks restart.
                before = page.get_by_text("SEC retrieved", exact=False).inner_text()
                page.locator('[data-testid="stMain"]').evaluate("element => { element.scrollTop = 0; }")
                page.screenshot(path=str(output / f"{name}-research.png"), full_page=True)
                page.reload()
                expect(download_button).to_be_visible(timeout=20000)
                expect(page.get_by_role("dialog")).to_be_hidden()
                expect(page.get_by_text("SEC retrieved", exact=False)).not_to_have_text(before)
                assert "cik=0000000001" in page.url
                assert not page.get_by_label("Contact email", exact=True).count()
                # Separate tabs must not silently inherit this tab's contact.
                other_tab = context.new_page()
                other_tab.goto(base)
                expect(other_tab.get_by_role("dialog")).to_be_visible(timeout=15000)
                expect(other_tab.get_by_label("Contact email", exact=True)).to_have_value("")
                other_tab.close()
                assert not errors, "Browser JavaScript exception(s) occurred"
                assert not external, "Fixture browser made an unexpected external request"
                report["viewports"].append({"name": name, "width": width, "height": height,
                                           "automatic_download": True, "contact_reload": True,
                                           "twenty_periods": True, "page_overflow": False})
                context.close()
            browser.close()
    finally:
        server.terminate()
        try:
            server.wait(timeout=10)
        except subprocess.TimeoutExpired:
            server.kill()
            server.wait(timeout=10)
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    destination = args.output or Path(tempfile.mkdtemp(prefix="financial-historian-browser-"))
    destination.mkdir(parents=True, exist_ok=True)
    result = run(args.chromium, destination)
    print(json.dumps(result, indent=2))
