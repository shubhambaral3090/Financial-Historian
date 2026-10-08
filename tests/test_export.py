from dataclasses import replace
from io import BytesIO

import openpyxl
import pytest

from financial_historian import export
from financial_historian.export import build_workbook
from financial_historian.models import Company, CoverageRecord, Fact, Filing, StatementBundle, StatementRow


SOURCE = "https://www.sec.gov/Archives/edgar/data/1234/000000123425000001/report.htm"


def fact(tag="us-gaap:Revenue", label="Revenue", unit="USD", value=42_250_000,
         end="2022-12-31", statement="Profit & Loss", **kwargs):
    data = dict(tag=tag, label=label, unit=unit, value=value, start=f"{end[:4]}-01-01",
                end=end, accession="0000001234-25-000001", filed="2025-02-15",
                form="10-K", source_url=SOURCE, statement=statement,
                extraction="companyfacts",
                selection_reason="Latest available comparable annual filed representation; not a confirmed restatement.")
    data.update(kwargs)
    return Fact(**data)


def sample_bundle():
    first = fact()
    latest = fact(value=50_000_000, end="2024-12-31")
    eps = fact(tag="us-gaap:EarningsPerShareDiluted", label="Diluted earnings per share",
               unit="USD/shares", value=3.21)
    shares = fact(tag="us-gaap:WeightedAverageNumberOfDilutedSharesOutstanding",
                  label="Diluted weighted average shares", unit="shares", value=13_000_000)
    euro = fact(tag="sample:EuropeanRevenue", label="Revenue reported in EUR", unit="EUR", value=3_000_000)
    assets = fact(tag="us-gaap:Assets", label="Total assets", value=80_000_000,
                  statement="Balance Sheet", start=None)
    cash = fact(tag="us-gaap:NetCashProvidedByUsedInOperatingActivities",
                label="Net cash from operations", value=-5_250_000, statement="Cash Flow")
    bundle = StatementBundle(Company("0000001234", "Example Corporation", "EXMP"),
                             periods=["2024-12-31", "2022-12-31", "2023-12-31"],
                             generated_at="2026-10-07T00:00:00+00:00")
    bundle.statements["Profit & Loss"] = [
        StatementRow(first.tag, first.label, first.unit, {first.end: first, latest.end: latest}),
        StatementRow(eps.tag, eps.label, eps.unit, {eps.end: eps}),
        StatementRow(shares.tag, shares.label, shares.unit, {shares.end: shares}),
        StatementRow(euro.tag, euro.label, euro.unit, {euro.end: euro}),
    ]
    bundle.statements["Balance Sheet"] = [StatementRow(assets.tag, assets.label, assets.unit, {assets.end: assets})]
    bundle.statements["Cash Flow"] = [StatementRow(cash.tag, cash.label, cash.unit, {cash.end: cash})]
    bundle.raw_facts = [first, latest, eps, shares, euro, assets, cash]
    bundle.filings = [Filing(first.accession, "10-K", first.filed, "2024-12-31", "report.htm", SOURCE)]
    bundle.coverage = [
        CoverageRecord(first.accession, "2024-12-31", first.filed, SOURCE, "structured",
                       "Annual entity-level XBRL facts extracted."),
        CoverageRecord("0000001234-92-000001", "1991-12-31", "1992-02-15",
                       "https://www.sec.gov/Archives/edgar/data/1234/000000123492000001/report.txt",
                       "unsupported", "Older text filing: no supported machine-readable statements."),
    ]
    bundle.warnings = ["Annual filing discovery does not establish complete company lifetime coverage."]
    return bundle


def loaded(bundle):
    return openpyxl.load_workbook(BytesIO(build_workbook(bundle)), data_only=False)


def test_workbook_preserves_units_missing_values_and_chronological_periods():
    bundle = sample_bundle()
    workbook = loaded(bundle)
    assert workbook.sheetnames[:8] == ["Profit & Loss", "Balance Sheet", "Cash Flow",
                                    "Cash Flow Detail", "Data Sheet", "Source Data", "Sources & Coverage", "Read Me"]
    assert len(workbook["Read Me"]._images) == 1
    sheet = workbook["Profit & Loss"]
    assert [sheet.cell(4, c).value for c in range(2, 5)] == ["2022-12-31", "2023-12-31", "2024-12-31"]
    assert sheet["B5"].value == 42.25
    assert sheet["C5"].value is None
    assert sheet["D5"].value == 50
    assert sheet["A5"].value == "Revenue (USD million)"
    assert sheet["B6"].value == 3.21
    assert sheet["A6"].value.endswith("(USD/shares)")
    assert sheet["B7"].value == 13_000_000
    assert sheet["A7"].value.endswith("(shares)")
    assert sheet["B8"].value == 3
    assert sheet["A8"].value.endswith("(EUR million)")
    assert workbook["Cash Flow"]["B5"].value == -5.25
    assert workbook["Balance Sheet"]["B5"].value == 80
    assert sheet.freeze_panes == "B5"
    raw = workbook["Source Data"]
    assert raw["E5"].value == 42_250_000
    assert raw["D7"].value == "USD/shares"
    assert raw["E7"].value == 3.21
    assert "'Profit & Loss'!B5" in raw["N5"].value


def test_data_sheet_is_a_clean_three_statement_annual_matrix():
    workbook = loaded(sample_bundle())
    matrix = workbook["Data Sheet"]
    assert matrix.max_column == 4
    assert matrix.freeze_panes == "B6"
    assert matrix.print_title_rows == "$1:$3"
    headings = [row for row in matrix.iter_rows() if row[0].value == "Reported line item"]
    assert len(headings) == 3
    assert all([cell.value for cell in row[1:]] == ["2022-12-31", "2023-12-31", "2024-12-31"]
               for row in headings)
    assert [matrix[f"A{row}"].value for row in (4, 11, 15)] == ["Profit & Loss", "Balance Sheet", "Cash Flow"]
    assert matrix["B6"].value == workbook["Profit & Loss"]["B5"].value == 42.25
    assert matrix["C6"].value is None  # Missing is never a fabricated zero.
    assert matrix["D6"].value == workbook["Profit & Loss"]["D5"].value == 50
    assert matrix["B7"].value == 3.21
    assert matrix["B8"].value == 13_000_000
    assert matrix["B9"].value == 3
    assert matrix["B13"].value == 80
    assert matrix["C13"].value is None and matrix["D13"].value is None
    assert matrix["B17"].value == -5.25
    assert matrix["B6"].hyperlink.target == SOURCE
    assert matrix["B6"].comment.text == workbook["Profit & Loss"]["B5"].comment.text
    assert "20 fiscal periods" in matrix["A2"].value
    assert "Source Data" in matrix["A3"].value
    assert not any("SEC tag" in str(cell.value) for row in matrix for cell in row)


def test_source_data_references_match_all_matrix_and_statement_values():
    workbook = loaded(sample_bundle())
    raw = workbook["Source Data"]
    headers = [cell.value for cell in raw[4]]
    assert headers == [
        "Statement", "SEC tag / row identifier", "Reported label", "Reported unit", "Source/XBRL value (unscaled)",
        "Annual display period start", "Annual display period end", "Accession", "Filed", "Form", "Extraction",
        "Selection reason", "Source filing", "Displayed cells", "Statement presentation value (unscaled)",
        "Original source instant", "Row role", "Reported section", "Context", "Dimensions", "Decimals",
    ]
    matrix_references = 0
    for record in raw.iter_rows(min_row=5, values_only=True):
        if not record[13]:
            continue
        presented = record[14] if record[14] is not None else record[4]
        expected = presented / 1_000_000 if record[3] in export.CURRENCIES else presented
        for reference in record[13].split("; "):
            name, coordinate = reference.split("!")
            name = name.strip("'")
            cell = workbook[name][coordinate]
            assert cell.value == expected
            assert f"SEC tag: {record[1]}" in cell.comment.text
            assert record[7] in cell.comment.text
            assert cell.hyperlink.target == record[12]
            matrix_references += name == "Data Sheet"
    assert matrix_references == 7


def test_provenance_and_old_filing_gaps_survive_export():
    workbook = loaded(sample_bundle())
    value = workbook["Profit & Loss"]["B5"]
    assert value.hyperlink.target == SOURCE
    for expected in ["us-gaap:Revenue", "2022-01-01", "2022-12-31", "2025-02-15",
                     "0000001234-25-000001", "companyfacts", "Selection reason", SOURCE]:
        assert expected in value.comment.text
    assert "not necessarily a confirmed restatement" in value.comment.text
    sources = workbook["Sources & Coverage"]
    assert sources["H5"].hyperlink.target == SOURCE
    assert sources["D7"].value == "1991-12-31"
    assert sources["F7"].value == "unsupported"
    assert "Older text filing" in sources["G7"].value
    assert sources["H7"].hyperlink.target.endswith("report.txt")
    assert sources["G8"].value == sample_bundle().warnings[0]
    readme_text = "\n".join(str(c.value or "") for row in workbook["Read Me"] for c in row)
    assert "not a complete statement presentation" in readme_text
    assert "company-specific custom tags" in readme_text
    assert "No normalization" in readme_text


def test_20_period_presentation_keeps_all_66_raw_years_and_custom_tags():
    from financial_historian.service import limit_statement_history

    periods = [f"{year}-12-31" for year in range(1960, 2026)]
    values = {period: fact(tag="example:SpecialRevenue", end=period, value=1_000_000) for period in periods}
    bundle = StatementBundle(Company("0000001234", "Long History Company"), periods=list(reversed(periods)))
    bundle.statements["Profit & Loss"] = [StatementRow("example:SpecialRevenue", "Special company revenue", "USD", values)]
    bundle.raw_facts = list(values.values())
    bundle = limit_statement_history(bundle)
    workbook = loaded(bundle)
    sheet = workbook["Profit & Loss"]
    assert sheet.max_column == 21
    assert sheet.cell(4, 2).value == "2006-12-31"
    assert sheet.cell(4, 21).value == "2025-12-31"
    assert sheet.cell(5, 21).value == 1
    assert "example:SpecialRevenue" in sheet.cell(5, 21).comment.text
    assert workbook["Source Data"].max_row == 70
    assert workbook["Source Data"]["G5"].value == "1960-12-31"
    assert workbook["Source Data"]["N5"].value is None
    matrix = workbook["Data Sheet"]
    assert matrix.max_column == 21
    assert matrix["B5"].value == "2006-12-31"
    assert matrix["U5"].value == "2025-12-31"


def test_export_rejects_unbounded_bundle_instead_of_silently_differing_from_browser():
    bundle = StatementBundle(Company("1", "Example"), periods=[f"{year}-12-31" for year in range(2000, 2026)])
    with pytest.raises(ValueError, match="latest 20 fiscal-period selection"):
        build_workbook(bundle)


def test_inconsistent_periods_and_units_are_blank_and_preserved_raw():
    wrong_period = fact(end="2023-12-31")
    wrong_unit = fact(unit="EUR", end="2024-12-31")
    bundle = StatementBundle(Company("0000001234", "Example"), periods=["2022-12-31", "2024-12-31"])
    bundle.statements["Profit & Loss"] = [StatementRow("us-gaap:Revenue", "Revenue", "USD", {
        "2022-12-31": wrong_period, "2024-12-31": wrong_unit,
    })]
    workbook = loaded(bundle)
    sheet = workbook["Profit & Loss"]
    assert sheet["B5"].value is None
    assert sheet["C5"].value is None
    assert "differs from column period" in sheet["B5"].comment.text
    assert "units were not mixed" in sheet["C5"].comment.text
    assert workbook["Source Data"]["G5"].value == "2023-12-31"
    assert workbook["Source Data"]["D6"].value == "EUR"
    assert workbook["Source Data"]["N5"].value is None


def test_reported_zero_is_zero_and_other_reported_units_are_unscaled():
    zero = fact(value=0)
    ratio = fact(tag="example:DebtRatio", label="Reported debt ratio", unit="pure", value=0.2345)
    bundle = StatementBundle(Company("1", "Example"), periods=[zero.end])
    bundle.statements["Profit & Loss"] = [StatementRow(f.tag, f.label, f.unit, {f.end: f}) for f in [zero, ratio]]
    workbook = loaded(bundle)
    assert workbook["Profit & Loss"]["B5"].value == 0
    assert workbook["Profit & Loss"]["B6"].value == 0.2345
    assert workbook["Profit & Loss"]["A6"].value.endswith("(pure)")


def test_cash_payment_uses_visible_outflow_sign_and_retains_positive_xbrl_value():
    payment = fact(tag="us-gaap:PaymentsToAcquirePropertyPlantAndEquipment",
                   label="Purchases of property, plant and equipment", value=20_000_000_000,
                   presentation_value=-20_000_000_000, statement="Cash Flow", extraction="inline-xbrl")
    bundle = StatementBundle(Company("0000001234", "Example"), periods=[payment.end], raw_facts=[payment])
    bundle.statements["Cash Flow"] = [StatementRow(payment.tag, payment.label, payment.unit, {payment.end: payment})]
    workbook = loaded(bundle)
    cell = workbook["Cash Flow Detail"]["B5"]
    assert cell.value == -20_000
    assert "Source/XBRL value (unscaled): 20000000000" in cell.comment.text
    assert "Statement presentation value (unscaled): -20000000000" in cell.comment.text
    assert workbook["Source Data"]["E5"].value == 20_000_000_000
    assert workbook["Source Data"]["O5"].value == -20_000_000_000
    assert workbook["Source Data"]["O4"].value == "Statement presentation value (unscaled)"


def test_zero_presentation_value_is_used_without_falling_back_to_source_value():
    reported = fact(value=42_000_000, presentation_value=0)
    bundle = StatementBundle(Company("0000001234", "Example"), periods=[reported.end], raw_facts=[reported])
    bundle.statements["Profit & Loss"] = [StatementRow(reported.tag, reported.label, reported.unit, {reported.end: reported})]
    workbook = loaded(bundle)
    assert workbook["Profit & Loss"]["B5"].value == 0
    assert workbook["Source Data"]["E5"].value == 42_000_000
    assert workbook["Source Data"]["O5"].value == 0


def test_beginning_and_ending_cash_keep_same_sec_tag_and_original_instant_dates():
    cash_tag = "us-gaap:CashCashEquivalentsRestrictedCashAndRestrictedCashEquivalents"
    beginning = fact(tag=cash_tag, label="Cash and restricted cash, beginning of year",
                     value=15_000_000_000, end="2024-12-31", start="2024-01-01",
                     statement="Cash Flow", extraction="inline-xbrl", row_role="cash_begin",
                     source_instant="2023-12-31")
    ending = replace(beginning, label="Cash and restricted cash, end of year", value=18_000_000_000,
                     row_role="cash_end", source_instant="2024-12-31")
    bundle = StatementBundle(Company("0000001234", "Example"), periods=[beginning.end], raw_facts=[beginning, ending])
    bundle.statements["Cash Flow"] = [StatementRow(f.tag, f.label, f.unit, {f.end: f}) for f in [beginning, ending]]
    workbook = loaded(bundle)
    cash = workbook["Cash Flow"]
    assert cash["B5"].value == 15_000
    assert cash["B6"].value == 18_000
    assert cash_tag in cash["B5"].comment.text and cash_tag in cash["B6"].comment.text
    assert "Annual display period start: 2024-01-01" in cash["B5"].comment.text
    assert "Annual display period end: 2024-12-31" in cash["B5"].comment.text
    assert "Original instant context: 2023-12-31" in cash["B5"].comment.text
    assert "Original instant context: 2024-12-31" in cash["B6"].comment.text
    assert "Row role: cash_begin" in cash["B5"].comment.text
    assert "Row role: cash_end" in cash["B6"].comment.text
    raw = workbook["Source Data"]
    assert raw["F5"].value == raw["F6"].value == "2024-01-01"
    assert raw["G5"].value == raw["G6"].value == "2024-12-31"
    assert raw["P5"].value == "2023-12-31"
    assert raw["P6"].value == "2024-12-31"
    assert raw["Q5"].value == "cash_begin"
    assert raw["Q6"].value == "cash_end"
    assert raw["B5"].value == raw["B6"].value == cash_tag
    assert raw["G4"].value == "Annual display period end"
    assert raw.auto_filter.ref == "A4:U6"


def test_all_external_text_is_literal_even_when_it_looks_like_a_formula():
    injection = '=HYPERLINK("https://attacker.invalid","click")'
    sample = fact(tag=injection, label="+SUM(1,1)", selection_reason="@SUM(1,1)", source_url=injection)
    raw = replace(sample, statement="-SUM(1,1)", extraction="\t=1+1", source_instant="=NOW()", row_role="@role")
    bundle = StatementBundle(Company("=1+1", injection, "@EXMP"), periods=[sample.end], generated_at="=NOW()")
    bundle.statements["Profit & Loss"] = [StatementRow(sample.tag, sample.label, sample.unit, {sample.end: sample})]
    bundle.raw_facts = [raw]
    bundle.filings = [Filing(injection, "=1+1", "@date", "+date", "-file", injection)]
    bundle.coverage = [CoverageRecord(injection, "=date", "@date", injection, "+status", "-detail")]
    bundle.warnings = [injection]
    workbook = loaded(bundle)
    assert workbook["Profit & Loss"]["A5"].value == "+SUM(1,1)"
    assert injection in workbook["Profit & Loss"]["B5"].comment.text
    assert workbook["Source Data"]["M5"].value == injection
    assert workbook["Source Data"]["P5"].value == "=NOW()"
    assert workbook["Source Data"]["Q5"].value == "@role"
    assert workbook["Source Data"]["M5"].hyperlink is None
    assert workbook["Profit & Loss"]["B5"].hyperlink is None
    for sheet in workbook:
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)


@pytest.mark.parametrize("source", [
    "http://www.sec.gov/Archives/edgar/data/1234/report.htm",
    "https://www.sec.gov.attacker.invalid/Archives/edgar/data/1234/report.htm",
    "https://www.sec.gov@attacker.invalid/Archives/edgar/data/1234/report.htm",
    "https://www.sec.gov/Archives/edgar/data/../report.htm",
    "https://www.sec.gov/Archives/edgar/data/%2e%2e/report.htm",
    "https://www.sec.gov/unrecognized/report.htm",
    "https://www.sec.gov:8443/Archives/edgar/data/1234/report.htm",
    "javascript:alert(1)",
    "file:///etc/passwd",
])
def test_untrusted_source_urls_remain_plain_text(source):
    bundle = sample_bundle()
    selected = replace(bundle.statements["Profit & Loss"][0].values["2022-12-31"], source_url=source)
    bundle.statements["Profit & Loss"][0].values["2022-12-31"] = selected
    workbook = loaded(bundle)
    assert workbook["Profit & Loss"]["B5"].hyperlink is None
    assert source in workbook["Profit & Loss"]["B5"].comment.text
    raw_rows = list(workbook["Source Data"].iter_rows(min_row=5, values_only=False))
    source_cells = [row[12] for row in raw_rows if row[12].value == source]
    assert source_cells
    assert all(cell.hyperlink is None for cell in source_cells)


def test_dimensions_fail_clearly_without_silent_history_truncation(monkeypatch):
    monkeypatch.setattr(export, "MAX_COLUMNS", 5)
    with pytest.raises(ValueError, match="No data was silently dropped"):
        build_workbook(sample_bundle())


def test_unlisted_periods_are_rejected_instead_of_silently_dropped():
    bundle = sample_bundle()
    bundle.periods.remove("2022-12-31")
    with pytest.raises(ValueError, match="absent from the bundle"):
        build_workbook(bundle)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), True])
def test_invalid_numbers_are_rejected(value):
    bundle = sample_bundle()
    bundle.raw_facts.append(fact(value=value))
    with pytest.raises(ValueError, match="finite numeric values"):
        build_workbook(bundle)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), True, "-100"])
def test_invalid_presentation_numbers_are_rejected(value):
    bundle = sample_bundle()
    bundle.raw_facts.append(fact(presentation_value=value))
    with pytest.raises(ValueError, match="presentation values.*finite"):
        build_workbook(bundle)


@pytest.mark.parametrize("name", ["X" * 32768, "Company\x00Name"])
def test_invalid_excel_text_is_rejected_with_clear_error(name):
    bundle = sample_bundle()
    bundle.company = replace(bundle.company, name=name)
    with pytest.raises(ValueError, match="text field"):
        build_workbook(bundle)


def test_verified_aliases_share_one_row_without_weakening_tag_validation():
    earlier = fact(tag="us-gaap:SalesRevenueNet", end="2016-12-31")
    modern = fact(tag="us-gaap:RevenueFromContractWithCustomerExcludingAssessedTax", end="2024-12-31")
    unrelated = fact(tag="us-gaap:Assets", end="2023-12-31")
    row = StatementRow(modern.tag, "Revenue", "USD", {
        earlier.end: earlier, unrelated.end: unrelated, modern.end: modern,
    }, accepted_tags={earlier.tag})
    bundle = StatementBundle(Company("1", "Example"), periods=[earlier.end, unrelated.end, modern.end])
    bundle.statements["Profit & Loss"] = [row]
    workbook = loaded(bundle)
    sheet = workbook["Profit & Loss"]
    assert sheet["B5"].value == sheet["D5"].value == 42.25
    assert earlier.tag in sheet["B5"].comment.text
    assert modern.tag in sheet["D5"].comment.text
    assert sheet["C5"].value is None
    assert "differs from its row tag" in sheet["C5"].comment.text
    assert workbook["Source Data"]["B6"].value == unrelated.tag


def test_statement_layout_groups_rows_bolds_totals_and_keeps_units_explicit():
    assets = fact(tag="us-gaap:Assets", label="Total assets", statement="Balance Sheet", section="Assets")
    debt = fact(tag="us-gaap:LongTermDebtNoncurrent", label="Term debt", statement="Balance Sheet", section="Noncurrent liabilities")
    current_debt = fact(tag="us-gaap:LongTermDebtCurrent", label="Term debt", statement="Balance Sheet", section="Current liabilities")
    bundle = StatementBundle(Company("1", "Example"), periods=[assets.end])
    bundle.statements["Balance Sheet"] = [
        StatementRow(assets.tag, assets.label, assets.unit, {assets.end: assets}, section=assets.section, is_total=True),
        StatementRow(current_debt.tag, current_debt.label, current_debt.unit, {current_debt.end: current_debt}, section=current_debt.section),
        StatementRow(debt.tag, debt.label, debt.unit, {debt.end: debt}, section=debt.section),
    ]
    workbook = loaded(bundle)
    sheet = workbook["Balance Sheet"]
    assert [sheet[f"A{r}"].value for r in range(5, 11)] == [
        "Assets", "Total assets", "Current liabilities", "Term debt", "Noncurrent liabilities", "Term debt",
    ]
    assert sheet["A6"].font.bold and sheet["B6"].font.bold
    assert not sheet["B8"].font.bold
    assert "USD million" in sheet["A2"].value
    assert sheet.max_column == 2
    assert "SEC tag" not in str([cell.value for cell in sheet[4]])
    assert workbook["Source Data"]["R5"].value == "Assets"
    assert "'Balance Sheet'!B6" in workbook["Source Data"]["N5"].value
    assert "'Data Sheet'!" in workbook["Source Data"]["N5"].value


def test_missing_history_keeps_the_same_columns_as_the_browser_and_explains_gaps():
    bundle = sample_bundle()
    bundle.periods = ["1960-12-31", *bundle.periods, "2025-12-31"]
    workbook = loaded(bundle)
    sheet = workbook["Profit & Loss"]
    selected = ["1960-12-31", "2022-12-31", "2023-12-31", "2024-12-31", "2025-12-31"]
    assert [cell.value for cell in sheet[4]][1:] == selected
    assert sheet["B5"].value is None and sheet["F5"].value is None
    assert sheet["D5"].value is None  # Keep the internal annual gap visible.
    for name in ("Balance Sheet", "Cash Flow", "Cash Flow Detail"):
        assert [cell.value for cell in workbook[name][4]][1:] == selected
    assert [cell.value for cell in workbook["Data Sheet"][5]][1:] == selected
    unavailable = [row for row in workbook["Sources & Coverage"].iter_rows(min_row=5, values_only=True)
                   if row[0] == "Unavailable period"]
    assert {row[3] for row in unavailable if row[6].startswith("Profit & Loss:")} == {"1960-12-31", "2023-12-31", "2025-12-31"}
    assert all(row[5] == "No selected values" for row in unavailable)
    assert all("matching the online presentation" in row[6] for row in unavailable)


def test_cash_overview_retains_all_company_detail_and_both_cell_references():
    totals = [fact(tag=f"us-gaap:NetCashProvidedByUsedIn{activity}Activities", label=f"Cash from {activity.lower()} activities",
                   statement="Cash Flow", value=value, section=f"{activity} activities")
              for activity, value in [("Operating", 80_000_000), ("Investing", -50_000_000), ("Financing", -10_000_000)]]
    custom = fact(tag="example:SpecialProceeds", label="Proceeds from a company-specific transaction",
                  statement="Cash Flow", value=17_000_000, section="Investing activities")
    bundle = StatementBundle(Company("1", "Example"), periods=[custom.end], raw_facts=[*totals, custom])
    bundle.statements["Cash Flow"] = [StatementRow(f.tag, f.label, f.unit, {f.end: f}, section=f.section, is_total=f != custom)
                                       for f in [*totals, custom]]
    workbook = loaded(bundle)
    main = workbook["Cash Flow"]
    assert [main[f"B{r}"].value for r in range(5, 8)] == [80, -50, -10]
    assert all(main[f"B{r}"].font.bold for r in range(5, 8))
    detail = workbook["Cash Flow Detail"]
    custom_cell = next(detail.cell(r, 2) for r in range(5, detail.max_row + 1)
                       if detail.cell(r, 1).value == custom.label)
    assert custom_cell.value == 17 and custom.tag in custom_cell.comment.text
    raw = workbook["Source Data"]
    assert "'Cash Flow'!B5" in raw["N5"].value
    assert "'Cash Flow Detail'!B6" in raw["N5"].value
    assert f"'Cash Flow Detail'!{custom_cell.coordinate}" in raw["N8"].value
    assert "'Data Sheet'!" in raw["N8"].value
    assert custom_cell.hyperlink.target == SOURCE


def test_export_reports_real_stages_and_supports_a_coverage_only_bundle():
    bundle = StatementBundle(Company("1", "Example"), periods=["1960-12-31"])
    messages = []
    workbook = openpyxl.load_workbook(BytesIO(build_workbook(bundle, progress=messages.append)))
    assert messages[0] == "Checking workbook values and source records…"
    assert messages[-2:] == ["Saving the Excel file…", "Excel file is ready."]
    assert any("source observations" in message for message in messages)
    for name in ("Profit & Loss", "Balance Sheet", "Cash Flow"):
        assert "No supported statement values" in workbook[name]["A3"].value
        assert workbook[name]["B4"].value == "1960-12-31"
    rows = list(workbook["Sources & Coverage"].iter_rows(min_row=5, values_only=True))
    gaps = [row for row in rows if row[0] == "Unavailable period"]
    assert len(gaps) == 3 and all(row[3] == "1960-12-31" for row in gaps)
    assert any(row[0] == "Presentation policy" and "20" in row[5] for row in rows)


def test_continuing_cash_flows_are_explicitly_separate_in_overview():
    total = fact(tag="us-gaap:NetCashProvidedByUsedInOperatingActivities", label="Cash from operating activities",
                 statement="Cash Flow", value=100_000_000)
    continuing = fact(tag="us-gaap:NetCashProvidedByUsedInOperatingActivitiesContinuingOperations",
                      label="Operating cash flow — continuing operations", statement="Cash Flow", value=90_000_000)
    bundle = StatementBundle(Company("1", "Example"), periods=[total.end])
    bundle.statements["Cash Flow"] = [StatementRow(f.tag, f.label, f.unit, {f.end: f}) for f in [total, continuing]]
    workbook = loaded(bundle)
    assert workbook["Cash Flow"]["B5"].value == 100
    assert workbook["Cash Flow"]["B6"].value == 90
    assert "continuing operations" in workbook["Cash Flow"]["A6"].value
    assert total.tag in workbook["Cash Flow"]["B5"].comment.text
    assert continuing.tag in workbook["Cash Flow"]["B6"].comment.text
    notes = dict(workbook["Read Me"].iter_rows(min_row=5, values_only=True))
    assert "not uniformly adjusted" in notes["EPS and stock splits"]


@pytest.mark.parametrize("statement, sheet_name, section, known_label, old_label", [
    ("Balance Sheet", "Balance Sheet", "Equity", "Total shareholders’ equity", "Legacy investment assets"),
    ("Cash Flow", "Cash Flow Detail", "Financing activities", "Financing proceeds", "Legacy cash movement"),
])
def test_unclassified_older_rows_cannot_visually_inherit_the_prior_section(
        statement, sheet_name, section, known_label, old_label):
    known = fact(tag="example:KnownLine", label=known_label, statement=statement, section=section)
    older = fact(tag="html:legacy_line", label=old_label, statement=statement, section="")
    bundle = StatementBundle(Company("1", "Example"), periods=[known.end])
    bundle.statements[statement] = [
        StatementRow(known.tag, known.label, known.unit, {known.end: known}, section=section),
        StatementRow(older.tag, older.label, older.unit, {older.end: older}),
    ]
    workbook = loaded(bundle)
    sheet = workbook[sheet_name]
    assert [sheet[f"A{r}"].value for r in range(5, 9)] == [
        section, known_label, "Other reported line items", old_label,
    ]
    assert sheet["A7"].font.bold and sheet["B7"].value is None
    assert sheet["B8"].value == 42.25
    assert "Reported section: Not supplied" in sheet["B8"].comment.text
    assert workbook["Source Data"]["R6"].value is None  # Preserve the missing source section.
    assert f"'{sheet_name}'!B8" in workbook["Source Data"]["N6"].value
