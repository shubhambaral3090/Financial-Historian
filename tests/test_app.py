"""Offline automatic UI workflows, privacy, and display/export consistency."""
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

import financial_historian.contact
import financial_historian.market
import financial_historian.service
import financial_historian.universe
from ui_fixtures import CONTACT, FixtureHistorian, FixtureMarket, FixtureUniverse


@pytest.fixture
def ui_file(tmp_path, monkeypatch):
    """Copied entry point isolates filesystem writes; all providers are offline."""
    monkeypatch.delenv("SEC_USER_AGENT", raising=False)
    monkeypatch.setattr(financial_historian.service, "FinancialHistorian", FixtureHistorian)
    monkeypatch.setattr(financial_historian.market, "YahooFinanceProvider", FixtureMarket)
    monkeypatch.setattr(financial_historian.universe, "UniverseProvider", FixtureUniverse)
    monkeypatch.setattr(financial_historian.contact, "browser_contact",
                        lambda contact=None: {"ready": True, "contact": None, "storageAvailable": True})
    for name in ("fail_search", "fail_resolve", "fail_generate", "empty"):
        monkeypatch.setattr(FixtureHistorian, name, False)
    for name, value in (("periods", 1), ("generate_calls", 0), ("resolve_calls", 0), ("search_queries", [])):
        monkeypatch.setattr(FixtureHistorian, name, value)
    monkeypatch.setattr(FixtureMarket, "fail", False)
    monkeypatch.setattr(FixtureMarket, "points", 2)
    monkeypatch.setattr(FixtureMarket, "calls", [])
    monkeypatch.setattr(FixtureUniverse, "fail", False)
    monkeypatch.setattr(FixtureUniverse, "calls", 0)
    st.cache_resource.clear()
    file = tmp_path / "app.py"
    file.write_text((Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8"), encoding="utf-8")
    yield file
    st.cache_resource.clear()


def button(app, label):
    return next(item for item in app.button if item.label == label)


def input_field(app, label):
    return next(item for item in app.text_input if item.label == label)


def choose_fixture(app, query="FIX"):
    input_field(app, "Company name, ticker, or CIK").input(query)
    button(app, "Search companies").click().run()
    next(item for item in app.button if item.label.startswith("Fictional UI Fixture")).click().run()
    return app


def preview_html(app):
    return " ".join(item.proto.body for item in app.get("html"))


def test_first_run_requires_contact_before_sec_requests(ui_file):
    app = AppTest.from_file(str(ui_file)).run()
    assert not app.exception
    assert input_field(app, "Your name or organization")
    assert input_field(app, "Contact email")
    assert not any(item.label == "Search companies" for item in app.button)
    assert FixtureHistorian.generate_calls == FixtureHistorian.resolve_calls == 0
    assert any("automated-access policy" in item.value for item in app.markdown)


@pytest.mark.parametrize("query", ["FIX", "Fictional UI Fixture", "0000000001"])
def test_selection_automatically_prepares_matching_workbook(ui_file, monkeypatch, query):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run(), query=query)
    assert not app.exception
    assert FixtureHistorian.search_queries == [query]
    assert FixtureHistorian.resolve_calls == FixtureHistorian.generate_calls == 1
    assert ("FIX", "1y", True) in FixtureMarket.calls
    workbook = app.session_state["workbook"]
    assert workbook.startswith(b"PK")
    assert app.session_state["workbook_name"] == "FIX_Financial_History.xlsx"
    assert any(item.label == "Download Excel workbook" for item in app.get("download_button"))
    labels = [item.label for item in app.button]
    assert not any(label in labels for label in ("Prepare Excel download", "Refresh market data", "Refresh this company", "Refresh company universe"))
    assert any(item.value == "Your workbook is ready" for item in app.subheader)
    assert "(1.2)" in preview_html(app) and "USD million" in preview_html(app)
    assert "2024-12-31" in preview_html(app)
    wb = load_workbook(BytesIO(workbook))
    cells = [cell for row in wb["Cash Flow Detail"] for cell in row if cell.value == -1.2]
    assert cells and all(cell.comment for cell in cells)
    assert "2024-12-31" in cells[0].comment.text
    assert "https://www.sec.gov/Archives/" in cells[0].comment.text
    assert "https://www.sec.gov/Archives/" in preview_html(app)
    for secret in (CONTACT, "Offline Fixture", "offline@example.invalid"):
        with ZipFile(BytesIO(workbook)) as archive:
            assert not any(secret.encode() in archive.read(name) for name in archive.namelist())
        assert secret not in "\n".join(app.session_state["diagnostics"])
    assert app.query_params["cik"] == ["0000000001"]
    app.run()
    assert not app.exception
    assert FixtureHistorian.generate_calls == 1, "widget reruns must not refetch filings"
    assert app.session_state["workbook"] == workbook


def test_suggestion_uses_same_sec_validated_automatic_page(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    app = AppTest.from_file(str(ui_file)).run()
    button(app, "FIX · Fictional UI Fixture").click().run()
    assert not app.exception
    assert FixtureHistorian.resolve_calls == FixtureHistorian.generate_calls == 1
    assert app.session_state["selected_company"].cik == "0000000001"
    assert app.session_state["workbook"]


def test_browser_reload_query_rechecks_both_providers(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    prior_time = app.session_state["research"].sec_retrieved_at
    refreshed = AppTest.from_file(str(ui_file))
    refreshed.query_params.update(app.query_params)
    refreshed.run()
    assert not refreshed.exception
    assert FixtureHistorian.generate_calls == 2
    assert sum(period == "1y" and force for _, period, force in FixtureMarket.calls) == 2
    assert refreshed.session_state["research"].sec_retrieved_at != prior_time
    assert refreshed.session_state["workbook"]


def test_failed_search_hides_private_provider_exception(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureHistorian, "fail_search", True)
    app = AppTest.from_file(str(ui_file)).run()
    input_field(app, "Company name, ticker, or CIK").input("FIX")
    button(app, "Search companies").click().run()
    assert not app.exception
    error = " ".join(item.value for item in app.error)
    assert "SEC company search is unavailable" in error
    assert "Private fixture" not in error
    assert CONTACT not in error + "\n".join(app.session_state["diagnostics"])


@pytest.mark.parametrize("settings", [json.dumps({"existing_option": "preserve"}), "broken-settings"])
def test_contact_is_session_only_and_ignores_legacy_settings(ui_file, settings):
    path = ui_file.parent / "data" / "settings.json"
    path.parent.mkdir()
    path.write_text(settings, encoding="utf-8")
    app = AppTest.from_file(str(ui_file)).run()
    input_field(app, "Your name or organization").input("Offline Fixture")
    input_field(app, "Contact email").input("offline@example.invalid")
    button(app, "Continue to Financial Historian").click().run()
    assert not app.exception
    assert path.read_text(encoding="utf-8") == settings
    assert app.session_state["contact_email"] == "offline@example.invalid"
    assert not button(app, "Search companies").disabled
    # AppTest retains old dialog-fragment nodes; actual DOM closure is checked
    # by browser_smoke.py. The sidebar must never contain contact inputs.
    assert not any(item.label == "Contact email" for item in app.sidebar.text_input)
    assert "offline@example.invalid" not in "\n".join(app.session_state["diagnostics"])


@pytest.mark.parametrize("name,email,message", [
    ("x" * 120, "y" * 120 + "@example.invalid", "250 characters"),
    ("東京", "offline@example.invalid", "HTTP header"),
])
def test_invalid_contact_rejected_without_request_or_write(ui_file, name, email, message):
    app = AppTest.from_file(str(ui_file)).run()
    input_field(app, "Your name or organization").input(name)
    input_field(app, "Contact email").input(email)
    button(app, "Continue to Financial Historian").click().run()
    assert not app.exception
    assert any(message in item.value for item in app.error)
    assert not (ui_file.parent / "data" / "settings.json").exists()
    assert FixtureHistorian.resolve_calls == 0


def test_market_failure_keeps_sec_and_matching_excel(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureMarket, "fail", True)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    assert app.session_state["workbook"]
    assert any("Yahoo Finance market data is unavailable" in item.value for item in app.warning)
    assert "Private technical" not in " ".join(item.value for item in app.warning)


@pytest.mark.parametrize("failure", ["fail_resolve", "fail_generate"])
def test_sec_failure_keeps_market_but_no_partial_workbook(ui_file, monkeypatch, failure):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureHistorian, failure, True)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    assert "workbook" not in app.session_state and "bundle" not in app.session_state
    assert any("102.00" in item.value for item in app.metric)
    assert any("SEC financial information could not be refreshed" in item.value for item in app.warning)
    visible = " ".join(item.value for item in app.warning) + "\n".join(app.session_state["diagnostics"])
    assert CONTACT not in visible and "Private fixture" not in visible


def test_universe_failure_does_not_block_direct_search(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureUniverse, "fail", True)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    assert app.session_state["workbook"]
    assert any("Direct SEC search still works" in item.value for item in app.warning)


def test_twenty_period_preview_and_workbook_keep_all_raw_sources(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureHistorian, "periods", 25)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    bundle = app.session_state["bundle"]
    assert sorted(bundle.periods) == [f"{year}-12-31" for year in range(2005, 2025)]
    assert len(bundle.raw_facts) == 25
    preview = preview_html(app)
    assert "2004-12-31" not in preview and "2005-12-31" in preview and "2024-12-31" in preview
    wb = load_workbook(BytesIO(app.session_state["workbook"]))
    values = {cell.value for row in wb["Cash Flow Detail"] for cell in row if isinstance(cell.value, str)}
    assert "2004-12-31" not in values
    assert "2005-12-31" in values and "2024-12-31" in values


def test_coverage_only_download_is_explicit(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    monkeypatch.setattr(FixtureHistorian, "empty", True)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    assert app.session_state["workbook"].startswith(b"PK")
    assert any("No annual financial values" in item.value for item in app.warning)
    assert next(item for item in app.metric if item.label == "Periods with figures").value == "0"
    assert any("1 discovered annual period has no extracted figures" in item.value for item in app.info)


def test_new_search_clears_previous_company_and_quote(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    input_field(app, "Company name, ticker, or CIK").input("NONE")
    button(app, "Search companies").click().run()
    assert not app.exception
    for key in ("workbook", "bundle", "selected_company", "research", "market_quote"):
        assert key not in app.session_state
    assert not any(item.label == "Download Excel workbook" for item in app.get("download_button"))


def test_chart_failure_does_not_hide_price_or_excel(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    def broken_chart(*args, **kwargs):
        raise ValueError("Offline chart-rendering fixture")
    monkeypatch.setattr(st, "line_chart", broken_chart)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    assert app.session_state["workbook"]
    assert any("chart could not be displayed" in item.value for item in app.info)
    assert any("102.00" in item.value for item in app.metric)


def test_statement_preview_escapes_reported_labels(ui_file, monkeypatch):
    monkeypatch.setenv("SEC_USER_AGENT", CONTACT)
    original_generate = FixtureHistorian.generate
    def special_labels(self, company, progress=None):
        bundle = original_generate(self, company, progress)
        row = bundle.statements["Cash Flow"][0]
        row.label = "<script>alert('fixture')</script> & capital expenditures"
        row.section = "Investing <activities>"
        row.is_total = True
        return bundle
    monkeypatch.setattr(FixtureHistorian, "generate", special_labels)
    app = choose_fixture(AppTest.from_file(str(ui_file)).run())
    assert not app.exception
    preview = preview_html(app)
    assert "<script>" not in preview
    assert "&lt;script&gt;" in preview and "Investing &lt;activities&gt;" in preview
    assert "statement-total" in preview
