from dataclasses import replace
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
import pytest

from financial_historian.analysis import build_analysis, Calculator
from financial_historian.models import Company, Fact, StatementBundle, StatementRow
from financial_historian.export import build_workbook


def historical_fixture():
    """Hand-designed data, not calculated from the implementation under test."""
    bundle = StatementBundle(Company("0000000001","Historical teaching fixture","HIST","3571"))
    mapping = {
        "Profit & Loss": {
            "Revenues": [800,900,1000,1200], "CostOfRevenue": [500,550,600,700],
            "GrossProfit": [300,350,400,500], "OperatingIncomeLoss": [80,90,100,180],
            "ProfitLoss": [60,70,80,120], "NetIncomeLoss": [58,68,78,118],
            "NetIncomeLossAvailableToCommonStockholdersBasic": [58,68,78,118],
            "InterestExpense": [10,10,10,12],
        },
        "Balance Sheet": {
            "Assets": [800,900,1100,1300], "LiabilitiesAndStockholdersEquity": [800,900,1100,1300],
            "AssetsCurrent": [250,300,350,400], "LiabilitiesCurrent": [100,120,140,160],
            "CashAndCashEquivalentsAtCarryingValue": [40,45,50,60],
            "ShortTermInvestments": [10,12,15,20], "AccountsReceivableNetCurrent": [70,80,90,110],
            "InventoryNet": [100,110,120,160], "AccountsPayableCurrent": [50,60,70,90],
            "CommonStockholdersEquity": [400,450,500,600],
            "DebtAndCapitalLeaseObligations": [100,120,150,180],
        },
        "Cash Flow": {
            "NetCashProvidedByUsedInOperatingActivities": [90,100,110,150],
            "PaymentsToAcquirePropertyPlantAndEquipment": [30,40,45,50],
        },
    }
    for statement,fields in mapping.items():
        for tag,values in fields.items():
            row = StatementRow("us-gaap:"+tag,tag,"USD")
            for year,v in zip(range(2021,2025),values):
                end = f"{year}-12-31"
                accession = f"0000000001-{year%100:02}-000001"
                f = Fact(row.tag,tag,"USD",v*1000000,None if statement=="Balance Sheet" else f"{year}-01-01",end,accession,
                         f"{year+1}-02-01","10-K",f"https://www.sec.gov/Archives/edgar/data/1/{accession.replace('-','')}/annual.htm",statement,decimals="-3")
                row.values[end] = f
                bundle.raw_facts.append(f)
            bundle.statements[statement].append(row)
    bundle.periods = [f"{year}-12-31" for year in range(2021,2025)]
    return bundle


def cell(analysis,key,p="2024-12-31"):
    return next(r for r in analysis.rows if r.definition.key==key).cells[p]


@pytest.mark.parametrize("key,expected",[
    ("revenue_growth",.2),("gross_margin",500/1200),("operating_margin",.15),
    ("roa",.1),("roe",118/550),("current_ratio",2.5),("quick_ratio",190/160),
    ("net_debt",120000000),("debt_assets",180/1300),("interest_coverage",15),
    ("receivable_days",100/1200*366),("inventory_days",140/700*366),
    ("payable_days",80/700*366),("ocf_margin",.125),("cash_conversion",1.25),
    ("ocf_capex",100000000),("earnings_cash",-30000000),("accrual_assets",-.025),
    ("asset_turnover",1),("operating_return",.15),("return_change",.05),
    ("margin_contribution",.05),("turnover_contribution",0),("interaction",0),
])
def test_hand_calculated_historical_metrics(key,expected):
    result = cell(build_analysis(historical_fixture()),key)
    assert result.value is not None, result.reason
    assert float(result.value) == pytest.approx(expected)
    assert result.sources


def test_bridge_identity_and_no_forecast_periods():
    result = build_analysis(historical_fixture())
    assert sum(cell(result,k).value for k in ("margin_contribution","turnover_contribution","interaction"))==cell(result,"return_change").value
    assert result.periods==["2021-12-31","2022-12-31","2023-12-31","2024-12-31"]
    assert next(c for c in result.checks if c.period=="2024-12-31" and c.label=="Operating return bridge").status=="passed"
    assert any(c.status=="insufficient information" for c in result.checks)


def test_note_input_uses_latest_vintage_and_preserves_richer_inline_context():
    bundle = historical_fixture()
    row = next(r for r in bundle.statements["Balance Sheet"] if r.tag.endswith(":CommonStockholdersEquity"))
    original = row.values.pop("2024-12-31")
    api = replace(original, statement="Disclosure", extraction="companyfacts-note", decimals="")
    inline = replace(api, extraction="inline-note", context_id="common-equity", decimals="-6")
    older = replace(api, accession="older", filed="2024-12-31", value=590000000)
    bundle.disclosures = [older, api, inline]
    result = Calculator(bundle).get("common_equity", "2024-12-31")
    assert result.value == Decimal("600000000")
    assert result.fact.context_id == "common-equity"
    assert result.fact.decimals == "-6"


def mutate(bundle,tag,**changes):
    for rows in bundle.statements.values():
        for row in rows:
            if row.tag=="us-gaap:"+tag:
                row.values["2024-12-31"] = replace(row.values["2024-12-31"],**changes)
                if "unit" in changes:
                    row.unit = changes["unit"]


@pytest.mark.parametrize("change,reason",[
    ({"value":0},"Denominator"),({"value":-100},"Denominator"),
    ({"value":1},"precision"),({"unit":"EUR"},"currencies"),
    ({"accession":"different-vintage"},"vintages"),
])
def test_invalid_denominator_currency_and_vintage(change,reason):
    bundle=historical_fixture()
    mutate(bundle,"LiabilitiesCurrent",**change)
    item=cell(build_analysis(bundle),"current_ratio")
    assert item.value is None
    assert reason in item.reason


def test_missing_is_not_zero_and_financial_industry_is_explicit():
    bundle=historical_fixture()
    bundle.statements["Balance Sheet"]=[r for r in bundle.statements["Balance Sheet"] if "ShortTermInvestments" not in r.tag]
    assert cell(build_analysis(bundle),"quick_ratio").value is None
    bundle.company=replace(bundle.company,sic="6021")
    result=build_analysis(bundle)
    assert cell(result,"operating_margin").value is None
    assert "financial" in cell(result,"operating_margin").reason
    assert any(c.status=="not applicable" for c in result.checks)


def test_lag_outside_window_stays_sourced_but_not_displayed():
    bundle=historical_fixture()
    bundle.periods=bundle.periods[1:]
    for rows in bundle.statements.values():
        for row in rows:
            row.values.pop("2021-12-31")
    result=build_analysis(bundle)
    assert cell(result,"roa","2022-12-31").value==Decimal(70000000)/Decimal(850000000)
    assert "2021-12-31" not in result.periods
    assert any(f.end=="2021-12-31" for f in cell(result,"roa","2022-12-31").sources)


def test_no_annual_growth_across_gap():
    bundle=historical_fixture()
    bundle.periods.remove("2023-12-31")
    bundle.raw_facts=[f for f in bundle.raw_facts if f.end!="2023-12-31"]
    for rows in bundle.statements.values():
        for row in rows:
            row.values.pop("2023-12-31")
    assert cell(build_analysis(bundle),"revenue_growth").value is None


def test_excel_cached_values_match_ui_and_formulas_are_linked():
    bundle=historical_fixture()
    analysis=build_analysis(bundle)
    payload=build_workbook(bundle)
    formulas=load_workbook(BytesIO(payload),data_only=False)
    values=load_workbook(BytesIO(payload),data_only=True)
    sheet=values["Historical Ratios"]
    row=next(row for row in sheet if row[0].value and str(row[0].value).startswith("Operating margin ["))
    address=row[-1].coordinate
    assert sheet[address].value==float(cell(analysis,"operating_margin").value)
    assert formulas["Historical Ratios"][address].value.startswith("='Calculation Inputs'!E")
    assert "Calculation Inputs!" in formulas["Historical Ratios"][address].comment.text
    assert any(c.data_type=="f" for row in formulas["Calculation Inputs"] for c in row)
    assert all(c.value is not None for row in formulas["Calculation Inputs"] for c in row if c.data_type=="f")
    assert "Context" in [c.value for c in values["Source Data"][4]]


def test_source_formula_injection_stays_literal():
    bundle=historical_fixture()
    bundle.company=replace(bundle.company,name="=HYPERLINK(\"https://bad.invalid\")")
    payload=build_workbook(bundle)
    w=load_workbook(BytesIO(payload))
    for sheet in w:
        for row in sheet:
            for c in row:
                if c.data_type=="f":
                    assert "HYPERLINK" not in c.value
